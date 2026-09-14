#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Revisor de PDF com Comentários Classificados - PREVENT Engenharia e Montagens
==============================================================================
Permite abrir um PDF e usar uma barra de ferramentas de marcação (Selecionar,
Mão, Comentário, Texto livre, Caneta, Retângulo, Elipse, Linha, Seta, Nuvem,
Sublinhar, Tachado, Realçar, Carimbo), para marcar pontos ou trechos da página
com um comentário classificado, e exportar:
  1) O PDF com as marcações inseridas como anotações reais.
  2) Uma planilha Excel (PROJETO, DATA, CODIGO ERRO, QUANTIDADE, TIPO DE ERRO,
     RESPONSÁVEL PELO ERRO, RESPONSÁVEL PELA VERIFICAÇÃO, NÚMERO DO DOCUMENTO,
     STATUS, REVISADO POR, DATA DA REVISÃO, VERIFICAÇÃO DA CORREÇÃO,
     VERIFICADO POR (CORREÇÃO)) pronta para tabela dinâmica, com abas de
     resumo por código de erro e painel de status.

Requisitos (instalar uma vez):
    pip install pymupdf pillow openpyxl sv-ttk --break-system-packages

Como rodar:
    python revisor_pdf.py

NOTA SOBRE OPACIDADE (item 19 da Parte 6): o formato de anotação PDF só tem
UM valor de opacidade (/CA) por anotação — não existe opacidade separada
nativa para borda e preenchimento. Este programa SIMULA as duas opacidades
pedidas misturando (alpha blend) a cor escolhida com o branco na proporção
equivalente antes de aplicar como cor sólida (ver `cor_com_opacidade_simulada`
abaixo). Essa é a abordagem documentada também no LEIA-ME.txt.
"""

import os
import io
import json
import base64
import tempfile
import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, colorchooser

try:
    import sv_ttk
except ImportError:
    sv_ttk = None

import fitz  # PyMuPDF
from PIL import Image, ImageDraw, ImageFont, ImageTk

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter


# --------------------------------------------------------------------------- #
# Identificação do programa
# --------------------------------------------------------------------------- #
VERSAO = "1.0"
DATA_CRIACAO = "Agosto/2026"
DESENVOLVEDOR = "Rafael Castro"

# --------------------------------------------------------------------------- #
# Categorias de classificação (edite esta lista se precisar ajustar depois)
# --------------------------------------------------------------------------- #
CLASSIFICACOES = [
    "ERRO DE CONCEPÇÃO DE PROJETO",
    "FALTA/ERRO DE DADOS DE ENTRADA",
    "NÃO UTILIZAÇÃO DE NORMA DE REFERÊNCIA",
    "ERRO DE DIMENSIONAMENTO/CÁLCULOS",
    "FALHA NO LEVANTAMENTO DE CAMPO",
    "FALTA DE DADOS E DETALHES PARA CONSTRUÇÃO E MONTAGEM",
    "INTERFERÊNCIA DE CAMPO",
    "INTERFACE INTERNA ENTRE DISCIPLINAS",
    "NÃO ATENDIMENTO À PADRONIZAÇÃO DO CLIENTE",
    "ERRO DE ESPECIFICAÇÃO DE MATERIAIS E EQUIPAMENTOS",
    "ERRO DE QUANTITATIVOS E PESO",
    "NÃO ATENDIMENTO A ESPECIFICAÇÃO E ESCOPO DO PROJETO",
    "ERRO DE CORTES E VISTAS",
    "ERRO DE COTAS, ELEVAÇÕES E COORDENADAS",
    "ERRO/FALTA DE DADOS NO CARIMBO",
    "ERRO DE TÍTULO DO DOCUMENTO",
    "ERRO DEVIDO A FALTA DE INDICAÇÃO DE REFERÊNCIAS",
    "ERRO/FALTA DE NOTAS E ERROS DE SIMBOLOGIA",
    "ERRO DE PORTUGUÊS",
    "ERRO DEVIDO TRANSCRIÇÃO DE DADOS",
    "DESENHO FORA DE ESCALA",
    "NÃO CONSIDERAÇÃO DE MUDANÇAS E DEFINIÇÕES AO LONGO DO PROJETO",
    "ERRO NO ATENDIMENTO A COMENTÁRIOS",
]

# Código curto de cada classificação (C01..C23), derivado automaticamente da
# ordem de CLASSIFICACOES — não altere os textos da lista acima.
CODIGOS_ERRO = {classificacao: f"C{idx:02d}" for idx, classificacao in enumerate(CLASSIFICACOES, start=1)}

APP_TITLE = "Revisor de PDF - PREVENT Engenharia e Montagens"


def exibicao_de_classificacao(classificacao):
    codigo = CODIGOS_ERRO.get(classificacao, "")
    return f"{codigo} - {classificacao}" if codigo else classificacao


def classificacao_de_exibicao(texto_exibicao):
    for classificacao, codigo in CODIGOS_ERRO.items():
        if texto_exibicao == f"{codigo} - {classificacao}":
            return classificacao
    return texto_exibicao  # já era o texto puro (compatibilidade)


# --------------------------------------------------------------------------- #
# Status de revisão colaborativa
# --------------------------------------------------------------------------- #
STATUS_REVISAO = ["Pendente", "Corrigido", "Não corrigido", "Reincidente", "Correto"]

STATUS_TAGS = {
    "Pendente": "status_pendente",
    "Corrigido": "status_corrigido",
    "Não corrigido": "status_nao_corrigido",
    "Reincidente": "status_reincidente",
    "Correto": "status_correto",
}

# Cor do MARCADOR (círculo numerado) conforme o status — padrão PP-ENG-002:
# vermelho = pendência, azul = corrigido, verde = corrigido E verificado,
# amarelo = item correto (não é um apontamento).
COR_MARCADOR_PENDENCIA = "#E53935"
COR_MARCADOR_CORRIGIDO = "#0070C0"
COR_MARCADOR_VERIFICADO = "#00B050"
COR_MARCADOR_CORRETO = "#FFD700"


def cor_status(item_ou_status):
    """Aceita tanto o dict completo do item (preferível, pois considera
    verificacao_correcao) quanto apenas a string de status (compatibilidade)."""
    if isinstance(item_ou_status, dict):
        status = item_ou_status.get("status", "Pendente")
        verificado = item_ou_status.get("verificacao_correcao", False)
    else:
        status = item_ou_status
        verificado = False

    if status == "Correto":
        return COR_MARCADOR_CORRETO
    if status == "Corrigido":
        return COR_MARCADOR_VERIFICADO if verificado else COR_MARCADOR_CORRIGIDO
    return COR_MARCADOR_PENDENCIA  # Pendente, Não corrigido, Reincidente


# --------------------------------------------------------------------------- #
# Elipse azul de confirmação + X verde de validação (PP-ENG-002, item 5.1.4.3)
# --------------------------------------------------------------------------- #
MARGEM_CONFIRMACAO = 12  # pontos (PDF) / pixels (canvas) de margem ao redor do marcador
COR_AZUL_CONFIRMACAO = "#0070C0"
COR_VERDE_VALIDACAO = "#00B050"


def rect_confirmacao(c, margem=MARGEM_CONFIRMACAO):
    """Retorna um fitz.Rect ao redor da geometria real do item (coordenadas do
    PDF), usado tanto para a elipse azul quanto para as diagonais do X verde."""
    x0, y0, x1, y1 = bbox_do_elemento(c)
    return fitz.Rect(x0 - margem, y0 - margem, x1 + margem, y1 + margem)


# --------------------------------------------------------------------------- #
# Cor amarela = "item correto" (não abre o pop-up de classificação)
# --------------------------------------------------------------------------- #
COR_AMARELO_CORRETO = "#FFFF00"
TEXTO_ITEM_CORRETO = "Item correto — validado sem apontamento"


def cor_e_amarela(hex_cor, tolerancia=40):
    """Compara por proximidade de cor (distância euclidiana em RGB 0-255),
    não exige o hex exato pixel a pixel."""
    if not hex_cor:
        return False
    r, g, b = (v * 255 for v in hex_para_rgb01(hex_cor))
    distancia = ((r - 255) ** 2 + (g - 255) ** 2 + (b - 0) ** 2) ** 0.5
    return distancia <= tolerancia


# --------------------------------------------------------------------------- #
# Ferramentas de marcação da barra de ferramentas
# --------------------------------------------------------------------------- #
FERRAMENTAS = {
    "selecionar":  {"label": "Selecionar",  "interacao": "selecionar", "icone": "🖱"},
    "mao":         {"label": "Mão",         "interacao": "mao",        "icone": "✋"},
    "comentario":  {"label": "Comentário",  "interacao": "clique",     "icone": "💬"},
    "texto_livre": {"label": "Texto livre", "interacao": "clique",     "icone": "✎"},
    "caneta":      {"label": "Caneta",      "interacao": "caneta",     "icone": "🖊"},
    "retangulo":   {"label": "Retângulo",   "interacao": "arrasto",    "icone": "▭"},
    "elipse":      {"label": "Elipse",      "interacao": "arrasto",    "icone": "⬭"},
    "linha":       {"label": "Linha",       "interacao": "arrasto",    "icone": "╱"},
    "seta":        {"label": "Seta",        "interacao": "arrasto",    "icone": "↗"},
    "nuvem":       {"label": "Nuvem",       "interacao": "arrasto",    "icone": "☁"},
    "sublinhado":  {"label": "Sublinhar",   "interacao": "arrasto",    "icone": "Sub"},
    "tachado":     {"label": "Tachado",     "interacao": "arrasto",    "icone": "Tac"},
    "realce":      {"label": "Realçar",     "interacao": "arrasto",    "icone": "🖍"},
    "carimbo":     {"label": "Carimbo",     "interacao": "clique",     "icone": "🖼"},
}

CORES_REALCE = {
    "Amarelo": "#FFFF00",
    "Verde": "#7CFC00",
    "Rosa": "#FF69B4",
    "Azul claro": "#87CEFA",
}

# Ferramentas de forma que usam o painel de estilo genérico (cor da borda,
# preenchimento quando aplicável, espessura e tracejado)
FERRAMENTAS_COM_BORDA = ("retangulo", "elipse", "nuvem", "linha", "seta")
FERRAMENTAS_COM_PREENCHIMENTO = ("retangulo", "elipse", "nuvem")

# Ferramentas cujo pop-up de classificação pode ser pulado se a cor ativa for
# amarela (item "Correto"). Comentário/Texto livre entram por completude do
# pedido, mas hoje só disparam de fato se tiverem um controle de cor associado.
FERRAMENTAS_SUJEITAS_A_CORRETO = (
    "comentario", "texto_livre", "retangulo", "elipse", "linha", "seta",
    "nuvem", "sublinhado", "tachado", "realce", "caneta",
)

ESTILOS_TRACEJADO = ["Contínua", "Tracejada", "Pontilhada", "Traço-ponto"]

DASHES_PDF = {
    "Contínua": None,
    "Tracejada": [6, 3],
    "Pontilhada": [1, 2],
    "Traço-ponto": [6, 2, 1, 2],
}

DASHES_CANVAS = {
    "Contínua": None,
    "Tracejada": (6, 3),
    "Pontilhada": (1, 2),
    "Traço-ponto": (6, 2, 1, 2),
}


def espessura_efetiva(item):
    """Pontilhada usa traço mais fino por padrão para reforçar a aparência de
    ponto (senão fica parecendo um tracejado curto quando a linha é grossa)."""
    espessura = item.get("espessura", 2)
    if item.get("tracejado") == "Pontilhada":
        return min(espessura, 1.5)
    return espessura


# Carimbo de verificação: tamanho padrão no PDF (pontos) e posições dos
# campos de texto dentro da imagem gerada (ajuste fino conforme a imagem
# oficial da empresa configurada pelo usuário)
CARIMBO_TAMANHO_PADRAO = (220, 140)
CARIMBO_POS_VERIFICADOR = (20, 90)
CARIMBO_POS_DATA = (20, 120)
CARIMBO_FONTE_TAMANHO = 20

# Logo oficial da PREVENT: se este arquivo existir na mesma pasta do script,
# é usado automaticamente como imagem padrão do Carimbo (item 17 da Parte 6).
# A logo (retangular, com "PREVENT" + "ENGENHARIA E MONTAGENS") não inclui a
# moldura/título/campos — esses continuam desenhados por código ao redor dela.
NOME_ARQUIVO_LOGO_PADRAO = "logo_prevent.png"
CAMINHO_LOGO_PADRAO = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), NOME_ARQUIVO_LOGO_PADRAO
)
# Tamanho da marcação no PDF (pontos) quando o carimbo com a logo padrão é
# usado — mantém a mesma proporção (3:2) da tela desenhada em código
# (CARIMBO_LOGO_CANVAS abaixo), para não distorcer ao inserir no PDF.
CARIMBO_TAMANHO_COM_LOGO = (240, 160)

# Tela (em pixels) usada para desenhar por código o carimbo com a logo, no
# padrão "moldura azul + logo + título + VERIFICADOR/DATA sublinhados"
# (modelo oficial fornecido pela PREVENT).
CARIMBO_LOGO_CANVAS = (600, 400)
COR_CARIMBO_BORDA = "#4472C4"
COR_CARIMBO_TEXTO = "#1F6ABA"
COR_CARIMBO_VALOR = "#1A1A1A"

COR_BORDA_SUAVE = "#D0D0D0"


def nome_cor_realce(hex_cor):
    for nome, valor in CORES_REALCE.items():
        if valor.lower() == (hex_cor or "").lower():
            return nome
    return "Amarelo"


def hex_para_rgb01(hex_cor):
    """Converte '#RRGGBB' em uma tupla (r, g, b) com valores 0-1, como o PyMuPDF espera."""
    hex_cor = (hex_cor or "#000000").lstrip("#")
    if len(hex_cor) != 6:
        hex_cor = "000000"
    return tuple(int(hex_cor[i:i + 2], 16) / 255 for i in (0, 2, 4))


def hex_para_rgba255(hex_cor, alpha=255):
    """Converte '#RRGGBB' em uma tupla (r, g, b, a) com valores 0-255, como o Pillow espera."""
    hex_cor = (hex_cor or "#000000").lstrip("#")
    if len(hex_cor) != 6:
        hex_cor = "000000"
    return tuple(int(hex_cor[i:i + 2], 16) for i in (0, 2, 4)) + (alpha,)


def cor_com_opacidade_simulada(hex_cor, opacidade_pct, fundo="#FFFFFF"):
    """PDF não tem opacidade nativa separada para borda x preenchimento numa
    mesma anotação (só existe um /CA geral). Simula visualmente misturando
    (alpha blend) a cor escolhida com o branco na proporção equivalente à
    opacidade pedida, devolvendo uma tupla RGB 0-1 pronta para set_colors()."""
    r1, g1, b1 = hex_para_rgb01(hex_cor)
    r2, g2, b2 = hex_para_rgb01(fundo)
    a = max(0.0, min(1.0, (opacidade_pct if opacidade_pct is not None else 100) / 100))
    return (r1 * a + r2 * (1 - a), g1 * a + g2 * (1 - a), b1 * a + b2 * (1 - a))


def hex_com_opacidade_simulada(hex_cor, opacidade_pct, fundo="#FFFFFF"):
    r, g, b = cor_com_opacidade_simulada(hex_cor, opacidade_pct, fundo)
    return "#%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))


def rect_de_marcacao(c, minimo=3.0):
    """Normaliza as coordenadas x0,y0,x1,y1 de uma marcação em um fitz.Rect válido,
    garantindo uma largura/altura mínima (retângulos degenerados quebram algumas
    anotações de marcação de texto no PyMuPDF)."""
    x0, x1 = sorted((c["x0"], c["x1"]))
    y0, y1 = sorted((c["y0"], c["y1"]))
    if x1 - x0 < minimo:
        cx = (x0 + x1) / 2
        x0, x1 = cx - minimo / 2, cx + minimo / 2
    if y1 - y0 < minimo:
        cy = (y0 + y1) / 2
        y0, y1 = cy - minimo / 2, cy + minimo / 2
    return fitz.Rect(x0, y0, x1, y1)


def rect_texto_livre(x, y, texto, fontsize):
    """Estima um retângulo de texto a partir do ponto inicial e do conteúdo digitado."""
    linhas = texto.split("\n") or [""]
    largura = max((len(l) for l in linhas), default=1) * fontsize * 0.55 + 10
    altura = len(linhas) * fontsize * 1.35 + 8
    largura = max(largura, fontsize * 3)
    altura = max(altura, fontsize * 1.5)
    return fitz.Rect(x, y, x + largura, y + altura)


def bbox_do_elemento(c):
    """Retorna (x0, y0, x1, y1) em coordenadas reais do PDF: a bounding box do
    elemento, usada para seleção/hit-test, destaque de seleção e movimentação."""
    tipo = c.get("tipo", "comentario")
    if tipo == "comentario":
        r = 10
        return (c["x"] - r, c["y"] - r, c["x"] + r, c["y"] + r)
    if tipo == "texto_livre":
        rect = rect_texto_livre(c["x"], c["y"], c["texto"], c.get("fontsize", 12))
        return (rect.x0, rect.y0, rect.x1, rect.y1)
    if tipo == "caneta":
        xs = [p[0] for p in c["pontos"]]
        ys = [p[1] for p in c["pontos"]]
        return (min(xs), min(ys), max(xs), max(ys))
    x0, x1 = sorted((c["x0"], c["x1"]))
    y0, y1 = sorted((c["y0"], c["y1"]))
    return (x0, y0, x1, y1)


def hoje_str():
    return datetime.date.today().strftime("%d/%m/%Y")


def _fonte_pil(tamanho, negrito=False):
    nomes = ["segoeuib.ttf", "arialbd.ttf"] if negrito else ["segoeui.ttf", "arial.ttf"]
    for nome in nomes:
        try:
            return ImageFont.truetype(nome, tamanho)
        except Exception:
            continue
    return ImageFont.load_default()


def _dividir_data(data_texto):
    """Separa 'DD/MM/AAAA' em (dia, mes, ano). Aceita texto incompleto ou
    fora do padrão sem lançar erro — devolve o que der para aproveitar."""
    partes = (data_texto or "").split("/")
    partes += [""] * (3 - len(partes))
    return partes[0].strip(), partes[1].strip(), partes[2].strip()


def _texto_centralizado(draw, cx, y, texto, fonte, fill):
    """Desenha `texto` centralizado horizontalmente em torno de x=cx."""
    bbox = draw.textbbox((0, 0), texto, font=fonte)
    tw = bbox[2] - bbox[0]
    draw.text((cx - tw / 2, y), texto, font=fonte, fill=fill)
    return tw


def _desenhar_campo_sublinhado(draw, x0, x1, y_linha, valor, fonte, cor_valor):
    """Desenha uma linha de sublinhado entre x0 e x1 em y_linha, com `valor`
    centralizado logo acima dela (campo de preenchimento tipo formulário)."""
    draw.line([(x0, y_linha), (x1, y_linha)], fill=cor_valor, width=2)
    if valor:
        _texto_centralizado(draw, (x0 + x1) / 2, y_linha - 30, valor, fonte, cor_valor)


def carimbo_padrao_gerado(verificador="", data_texto=""):
    """Gera por código um carimbo simples (moldura azul arredondada + título
    + campos VERIFICADOR/DATA sublinhados), usado quando nenhuma imagem (nem
    a logo padrão) está disponível — mesmo layout do carimbo com logo, só
    que sem a faixa da logo no topo."""
    w, h = 560, 300
    img = Image.new("RGBA", (w, h), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    borda_rgba = hex_para_rgba255(COR_CARIMBO_BORDA)
    texto_rgba = hex_para_rgba255(COR_CARIMBO_TEXTO)
    valor_rgba = hex_para_rgba255(COR_CARIMBO_VALOR)

    draw.rounded_rectangle([8, 8, w - 8, h - 8], radius=36, outline=borda_rgba, width=9)

    fonte_titulo = _fonte_pil(30, negrito=True)
    _texto_centralizado(draw, w / 2, 34, "COPIA DE VERIFICAÇÃO", fonte_titulo, texto_rgba)

    fonte_label = _fonte_pil(22, negrito=False)
    fonte_valor = _fonte_pil(22, negrito=False)
    margem = 42

    y_verificador = 128
    draw.text((margem, y_verificador), "VERIFICADOR:", font=fonte_label, fill=texto_rgba)
    bbox = draw.textbbox((0, 0), "VERIFICADOR:", font=fonte_label)
    x_campo = margem + (bbox[2] - bbox[0]) + 10
    _desenhar_campo_sublinhado(
        draw, x_campo, w - margem, y_verificador + 34, verificador, fonte_valor, valor_rgba
    )

    y_data = 210
    dia, mes, ano = _dividir_data(data_texto)
    draw.text((margem, y_data), "DATA:", font=fonte_label, fill=texto_rgba)
    bbox = draw.textbbox((0, 0), "DATA:", font=fonte_label)
    x = margem + (bbox[2] - bbox[0]) + 14
    largura_dia, largura_mes, largura_ano = 70, 70, 100
    y_linha = y_data + 34

    _desenhar_campo_sublinhado(draw, x, x + largura_dia, y_linha, dia, fonte_valor, valor_rgba)
    x += largura_dia + 6
    draw.text((x, y_data), "/", font=fonte_label, fill=texto_rgba)
    x += 16

    _desenhar_campo_sublinhado(draw, x, x + largura_mes, y_linha, mes, fonte_valor, valor_rgba)
    x += largura_mes + 6
    draw.text((x, y_data), "/", font=fonte_label, fill=texto_rgba)
    x += 16

    _desenhar_campo_sublinhado(draw, x, x + largura_ano, y_linha, ano, fonte_valor, valor_rgba)

    return img


def carimbo_com_logo_gerado(caminho_logo, verificador="", data_texto=""):
    """Monta o carimbo oficial da PREVENT: moldura azul arredondada, logo da
    empresa centralizada no topo, título 'COPIA DE VERIFICAÇÃO' e os campos
    VERIFICADOR e DATA (dia/mês/ano) com linhas de sublinhado — no mesmo
    padrão do modelo oficial fornecido pela empresa."""
    w, h = CARIMBO_LOGO_CANVAS
    img = Image.new("RGBA", (w, h), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    borda_rgba = hex_para_rgba255(COR_CARIMBO_BORDA)
    texto_rgba = hex_para_rgba255(COR_CARIMBO_TEXTO)
    valor_rgba = hex_para_rgba255(COR_CARIMBO_VALOR)

    draw.rounded_rectangle([10, 10, w - 10, h - 10], radius=44, outline=borda_rgba, width=10)

    # ---- logo centralizada no topo ---- #
    logo = Image.open(caminho_logo).convert("RGBA")
    area_largura, area_altura = w - 140, 150
    escala = min(area_largura / logo.width, area_altura / logo.height)
    novo_tam = (max(1, int(logo.width * escala)), max(1, int(logo.height * escala)))
    logo = logo.resize(novo_tam, Image.LANCZOS)
    pos_x = (w - novo_tam[0]) // 2
    pos_y = 34
    img.alpha_composite(logo, (pos_x, pos_y))

    # ---- título ---- #
    y_titulo = pos_y + novo_tam[1] + 22
    fonte_titulo = _fonte_pil(34, negrito=True)
    _texto_centralizado(draw, w / 2, y_titulo, "COPIA DE VERIFICAÇÃO", fonte_titulo, texto_rgba)

    fonte_label = _fonte_pil(24, negrito=False)
    fonte_valor = _fonte_pil(24, negrito=False)
    margem = 46

    # ---- linha VERIFICADOR ---- #
    y_verificador = y_titulo + 62
    draw.text((margem, y_verificador), "VERIFICADOR:", font=fonte_label, fill=texto_rgba)
    bbox = draw.textbbox((0, 0), "VERIFICADOR:", font=fonte_label)
    x_campo = margem + (bbox[2] - bbox[0]) + 10
    _desenhar_campo_sublinhado(
        draw, x_campo, w - margem, y_verificador + 36, verificador, fonte_valor, valor_rgba
    )

    # ---- linha DATA (dia / mês / ano) ---- #
    y_data = y_verificador + 82
    dia, mes, ano = _dividir_data(data_texto)
    draw.text((margem, y_data), "DATA:", font=fonte_label, fill=texto_rgba)
    bbox = draw.textbbox((0, 0), "DATA:", font=fonte_label)
    x = margem + (bbox[2] - bbox[0]) + 16
    largura_dia, largura_mes, largura_ano = 76, 76, 108
    y_linha = y_data + 36

    _desenhar_campo_sublinhado(draw, x, x + largura_dia, y_linha, dia, fonte_valor, valor_rgba)
    x += largura_dia + 8
    draw.text((x, y_data), "/", font=fonte_label, fill=texto_rgba)
    x += 20

    _desenhar_campo_sublinhado(draw, x, x + largura_mes, y_linha, mes, fonte_valor, valor_rgba)
    x += largura_mes + 8
    draw.text((x, y_data), "/", font=fonte_label, fill=texto_rgba)
    x += 20

    _desenhar_campo_sublinhado(draw, x, x + largura_ano, y_linha, ano, fonte_valor, valor_rgba)

    return img


def gerar_imagem_carimbo(caminho_imagem, verificador, data_texto):
    """Combina a imagem de fundo do carimbo (imagem configurada manualmente,
    ou a logo_prevent.png padrão se existir, ou o carimbo gerado por código
    como último recurso) com os campos de texto variáveis (VERIFICADOR/DATA),
    e devolve bytes PNG prontos para inserir no PDF via page.insert_image().

    Quando a logo padrão ou o carimbo gerado por código são usados, o layout
    completo (título + campos sublinhados) já é desenhado internamente por
    `carimbo_com_logo_gerado`/`carimbo_padrao_gerado`. Só no caso de imagem
    própria configurada pelo usuário (sem esse layout embutido) os campos
    VERIFICADOR/DATA são sobrepostos numa posição fixa e configurável."""
    if caminho_imagem and os.path.exists(caminho_imagem):
        base = Image.open(caminho_imagem).convert("RGBA")
        draw = ImageDraw.Draw(base)
        fonte = _fonte_pil(CARIMBO_FONTE_TAMANHO)
        draw.text(
            CARIMBO_POS_VERIFICADOR, f"VERIFICADOR: {verificador}",
            font=fonte, fill=(0, 0, 0, 255),
        )
        draw.text(
            CARIMBO_POS_DATA, f"DATA: {data_texto}",
            font=fonte, fill=(0, 0, 0, 255),
        )
    elif os.path.exists(CAMINHO_LOGO_PADRAO):
        base = carimbo_com_logo_gerado(CAMINHO_LOGO_PADRAO, verificador, data_texto)
    else:
        base = carimbo_padrao_gerado(verificador, data_texto)

    buffer = io.BytesIO()
    base.save(buffer, format="PNG")
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Tooltip simples (Enter/Leave em qualquer widget)
# --------------------------------------------------------------------------- #
class Tooltip:
    def __init__(self, widget, texto):
        self.widget = widget
        self.texto = texto
        self.tip = None
        widget.bind("<Enter>", self._mostrar, add="+")
        widget.bind("<Leave>", self._esconder, add="+")

    def _mostrar(self, event=None):
        if self.tip or not self.texto:
            return
        try:
            x = self.widget.winfo_rootx() + 6
            y = self.widget.winfo_rooty() + self.widget.winfo_height() + 4
        except tk.TclError:
            return
        self.tip = tk.Toplevel(self.widget)
        self.tip.wm_overrideredirect(True)
        self.tip.wm_geometry(f"+{x}+{y}")
        tk.Label(
            self.tip, text=self.texto, background="#FFFFE0", relief="solid",
            borderwidth=1, font=("Segoe UI", 8), padx=4, pady=2,
        ).pack()

    def _esconder(self, event=None):
        if self.tip:
            self.tip.destroy()
            self.tip = None


# --------------------------------------------------------------------------- #
# Diálogo: dados gerais do documento
# --------------------------------------------------------------------------- #
class DadosDocumentoDialog(tk.Toplevel):
    def __init__(self, master, valores_atuais=None):
        super().__init__(master)
        self.title("Dados do documento")
        self.resizable(False, False)
        self.result = None
        self.grab_set()

        valores_atuais = valores_atuais or {}

        campos = [
            ("projeto", "Projeto:"),
            ("codigo_doc", "Número do documento:"),
            ("elaborado_por", "Elaborado por:"),
            ("data_elaboracao", "Data de elaboração (dd/mm/aaaa):"),
            ("verificado_por", "Verificado por:"),
            ("data_verificacao", "Data de verificação (dd/mm/aaaa):"),
        ]

        self.vars = {}
        for i, (chave, rotulo) in enumerate(campos):
            ttk.Label(self, text=rotulo, anchor="w").grid(
                row=i, column=0, sticky="w", padx=10, pady=6
            )
            var = tk.StringVar(value=valores_atuais.get(chave, ""))
            if chave == "data_elaboracao" and not var.get():
                var.set(hoje_str())
            entry = ttk.Entry(self, textvariable=var, width=35)
            entry.grid(row=i, column=1, padx=10, pady=6)
            self.vars[chave] = var

        btn_frame = ttk.Frame(self)
        btn_frame.grid(row=len(campos), column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame, text="Confirmar", width=12, command=self._confirmar).pack(
            side="left", padx=5
        )
        ttk.Button(btn_frame, text="Cancelar", width=12, command=self.destroy).pack(
            side="left", padx=5
        )

        self.wait_window(self)

    def _confirmar(self):
        self.result = {k: v.get().strip() for k, v in self.vars.items()}
        self.destroy()


# --------------------------------------------------------------------------- #
# Diálogo: inserir / editar comentário (comum a todas as ferramentas de marcação)
# --------------------------------------------------------------------------- #
class ComentarioDialog(tk.Toplevel):
    def __init__(self, master, pagina, tipo="comentario", texto="", classificacao="",
                 status=None, verificacao_correcao=False, fontsize=12, cor=None):
        super().__init__(master)
        self.tipo = tipo
        rotulo_tipo = FERRAMENTAS.get(tipo, {}).get("label", "Comentário")
        self.title(f"Comentário - {rotulo_tipo} - página {pagina + 1}")
        self.resizable(False, False)
        self.result = None
        self.grab_set()

        if cor is None:
            cor = "#FFFF00" if tipo == "realce" else "#000000"

        ttk.Label(self, text="Classificação:", anchor="w").grid(
            row=0, column=0, sticky="w", padx=10, pady=(10, 4)
        )
        self.classificacao_var = tk.StringVar(
            value=exibicao_de_classificacao(classificacao or CLASSIFICACOES[0])
        )
        combo = ttk.Combobox(
            self,
            textvariable=self.classificacao_var,
            values=[exibicao_de_classificacao(c) for c in CLASSIFICACOES],
            width=65,
            state="readonly",
        )
        combo.grid(row=1, column=0, padx=10, pady=(0, 10), sticky="w")

        ttk.Label(self, text="Comentário:", anchor="w").grid(
            row=2, column=0, sticky="w", padx=10
        )
        self.texto_widget = tk.Text(
            self, width=60, height=6, wrap="word",
            borderwidth=0, highlightthickness=1, highlightbackground=COR_BORDA_SUAVE,
        )
        self.texto_widget.grid(row=3, column=0, padx=10, pady=(0, 10))
        self.texto_widget.insert("1.0", texto)
        self.texto_widget.focus_set()

        ttk.Label(self, text="Status:", anchor="w").grid(row=4, column=0, sticky="w", padx=10)
        self.status_var = tk.StringVar(value=status or "Pendente")
        ttk.Combobox(
            self, textvariable=self.status_var, values=STATUS_REVISAO,
            state="readonly", width=20,
        ).grid(row=5, column=0, padx=10, pady=(0, 10), sticky="w")

        self.verificacao_var = tk.BooleanVar(value=verificacao_correcao)
        self.chk_verificacao = ttk.Checkbutton(
            self, text="Verificação da correção", variable=self.verificacao_var
        )
        self.chk_verificacao.grid(row=6, column=0, sticky="w", padx=10)
        ttk.Label(
            self, text="Disponível após marcar como Corrigido",
            foreground="#888888", font=("Segoe UI", 8),
        ).grid(row=7, column=0, sticky="w", padx=10, pady=(0, 10))

        def _atualizar_estado_verificacao(*_args):
            if self.status_var.get() == "Corrigido":
                self.chk_verificacao.state(["!disabled"])
            else:
                self.chk_verificacao.state(["disabled"])

        self.status_var.trace_add("write", _atualizar_estado_verificacao)
        _atualizar_estado_verificacao()

        linha_extra = 8
        self.fontsize_var = tk.IntVar(value=fontsize)
        self.cor_var = tk.StringVar(value=cor)
        self.combo_cor_realce = None
        self.swatch = None

        if tipo == "texto_livre":
            frame_extra = ttk.Frame(self)
            frame_extra.grid(row=linha_extra, column=0, sticky="w", padx=10, pady=(0, 10))
            ttk.Label(frame_extra, text="Tamanho da fonte:").pack(side="left")
            ttk.Spinbox(
                frame_extra, from_=6, to=72, width=5, textvariable=self.fontsize_var
            ).pack(side="left", padx=(4, 16))
            self._campo_cor(frame_extra, "Cor do texto:")
            linha_extra += 1
        elif tipo == "realce":
            frame_extra = ttk.Frame(self)
            frame_extra.grid(row=linha_extra, column=0, sticky="w", padx=10, pady=(0, 10))
            ttk.Label(frame_extra, text="Cor do realce:").pack(side="left")
            self.combo_cor_realce = ttk.Combobox(
                frame_extra,
                values=list(CORES_REALCE.keys()),
                state="readonly",
                width=12,
            )
            self.combo_cor_realce.set(nome_cor_realce(cor))
            self.combo_cor_realce.pack(side="left", padx=4)
            linha_extra += 1

        btn_frame = ttk.Frame(self)
        btn_frame.grid(row=linha_extra, column=0, pady=(0, 10))
        ttk.Button(btn_frame, text="Salvar", width=12, command=self._salvar).pack(
            side="left", padx=5
        )
        ttk.Button(btn_frame, text="Cancelar", width=12, command=self.destroy).pack(
            side="left", padx=5
        )

        self.wait_window(self)

    def _campo_cor(self, parent, rotulo):
        ttk.Label(parent, text=rotulo).pack(side="left")
        # Label tk "clássico" propositalmente mantido aqui: é a única forma simples
        # de mostrar uma amostra de cor arbitrária (ttk.Label não tem 'bg' próprio).
        self.swatch = tk.Label(parent, text="    ", bg=self.cor_var.get(), relief="solid", borderwidth=1)
        self.swatch.pack(side="left", padx=(4, 4))

        def escolher():
            escolha = colorchooser.askcolor(color=self.cor_var.get(), title=rotulo)
            if escolha and escolha[1]:
                self.cor_var.set(escolha[1])
                self.swatch.config(bg=escolha[1])

        ttk.Button(parent, text="Escolher...", command=escolher).pack(side="left")

    def _salvar(self):
        texto = self.texto_widget.get("1.0", "end").strip()
        if not texto:
            messagebox.showwarning("Aviso", "Digite o texto do comentário.")
            return
        resultado = {
            "texto": texto,
            "classificacao": classificacao_de_exibicao(self.classificacao_var.get()),
            "status": self.status_var.get(),
            "verificacao_correcao": (
                self.verificacao_var.get() if self.status_var.get() == "Corrigido" else False
            ),
        }
        if self.tipo == "texto_livre":
            resultado["fontsize"] = self.fontsize_var.get()
            resultado["cor"] = self.cor_var.get()
        elif self.tipo == "realce":
            resultado["cor"] = CORES_REALCE[self.combo_cor_realce.get()]
        self.result = resultado
        self.destroy()


# --------------------------------------------------------------------------- #
# Diálogo: carimbo de verificação (verificador + data)
# --------------------------------------------------------------------------- #
class CarimboDialog(tk.Toplevel):
    def __init__(self, master, verificador="", data=""):
        super().__init__(master)
        self.title("Carimbo de verificação")
        self.resizable(False, False)
        self.result = None
        self.grab_set()

        ttk.Label(self, text="Verificador:", anchor="w").grid(
            row=0, column=0, sticky="w", padx=10, pady=(10, 4)
        )
        self.verificador_var = tk.StringVar(value=verificador)
        ttk.Entry(self, textvariable=self.verificador_var, width=35).grid(
            row=1, column=0, padx=10, pady=(0, 10)
        )

        ttk.Label(self, text="Data (dd/mm/aaaa):", anchor="w").grid(
            row=2, column=0, sticky="w", padx=10
        )
        self.data_var = tk.StringVar(value=data or hoje_str())
        ttk.Entry(self, textvariable=self.data_var, width=35).grid(
            row=3, column=0, padx=10, pady=(0, 10)
        )

        btn_frame = ttk.Frame(self)
        btn_frame.grid(row=4, column=0, pady=(0, 10))
        ttk.Button(btn_frame, text="Inserir", width=12, command=self._confirmar).pack(
            side="left", padx=5
        )
        ttk.Button(btn_frame, text="Cancelar", width=12, command=self.destroy).pack(
            side="left", padx=5
        )

        self.wait_window(self)

    def _confirmar(self):
        self.result = {
            "verificador": self.verificador_var.get().strip(),
            "data": self.data_var.get().strip(),
        }
        self.destroy()


# --------------------------------------------------------------------------- #
# Diálogo: identificação do revisor (perguntado uma vez por sessão)
# --------------------------------------------------------------------------- #
class RevisorDialog(tk.Toplevel):
    def __init__(self, master, nome_atual=""):
        super().__init__(master)
        self.title("Quem está revisando agora?")
        self.resizable(False, False)
        self.result = None
        self.grab_set()

        ttk.Label(self, text="Nome do revisor:", anchor="w").grid(
            row=0, column=0, sticky="w", padx=10, pady=(10, 4)
        )
        self.nome_var = tk.StringVar(value=nome_atual)
        entry = ttk.Entry(self, textvariable=self.nome_var, width=35)
        entry.grid(row=1, column=0, padx=10, pady=(0, 10))
        entry.focus_set()

        btn_frame = ttk.Frame(self)
        btn_frame.grid(row=2, column=0, pady=(0, 10))
        ttk.Button(btn_frame, text="Confirmar", width=12, command=self._confirmar).pack(
            side="left", padx=5
        )
        ttk.Button(btn_frame, text="Cancelar", width=12, command=self.destroy).pack(
            side="left", padx=5
        )

        self.wait_window(self)

    def _confirmar(self):
        nome = self.nome_var.get().strip()
        if not nome:
            messagebox.showwarning("Aviso", "Digite o nome do revisor.")
            return
        self.result = nome
        self.destroy()


# --------------------------------------------------------------------------- #
# Diálogo: Sobre
# --------------------------------------------------------------------------- #
class SobreDialog(tk.Toplevel):
    def __init__(self, master):
        super().__init__(master)
        self.title("Sobre")
        self.resizable(False, False)
        self.grab_set()

        texto = (
            f"{APP_TITLE}\n\n"
            f"Versão: {VERSAO}\n"
            f"Criado em: {DATA_CRIACAO}\n"
            f"Desenvolvido por: {DESENVOLVEDOR}"
        )
        ttk.Label(self, text=texto, justify="center", padding=20).pack()
        ttk.Button(self, text="Fechar", command=self.destroy).pack(pady=(0, 15))

        self.update_idletasks()
        x = master.winfo_rootx() + max(0, (master.winfo_width() - self.winfo_width()) // 2)
        y = master.winfo_rooty() + max(0, (master.winfo_height() - self.winfo_height()) // 2)
        self.geometry(f"+{x}+{y}")


# --------------------------------------------------------------------------- #
# Aplicação principal
# --------------------------------------------------------------------------- #
class RevisorPDFApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("1200x760")

        self.doc = None
        self.pdf_path = None
        self.pagina_atual = 0
        self.zoom = 1.4
        self.comentarios = []  # lista de dicts (comentários e marcações visuais)
        self.dados_documento = {}
        self.tk_img = None

        self._arrasto_inicio = None
        self._preview_ids = []

        # modo Selecionar: mover elemento(s) e seleção por retângulo (marquee)
        self._mover_idx = None
        self._mover_inicio_px = None
        self._mover_original_grupo = {}
        self._marquee_inicio = None

        # ferramenta Caneta: pontos capturados em coordenadas de canvas
        self._pontos_caneta_canvas = []
        self._preview_caneta_id = None

        # revisão colaborativa
        self.revisor_atual = None
        self._avisou_pdf_embutido = False

        self._lateral_visivel = False

        self._montar_interface()
        self._atalhos()

    # ------------------------- interface ------------------------- #
    def _montar_interface(self):
        self._montar_menu()

        # Barra superior
        barra = ttk.Frame(self.root)
        barra.pack(side="top", fill="x", padx=6, pady=6)

        self._botao_icone(barra, "📂", "Abrir PDF", self.abrir_pdf)
        self._botao_icone(barra, "🗒", "Dados do documento", self.editar_dados_documento)
        self._botao_icone(barra, "💾", "Salvar projeto", self.salvar_projeto)
        self._botao_icone(barra, "📁", "Abrir projeto", self.abrir_projeto)
        self._botao_icone(barra, "👤", "Identificar revisor", self._identificar_revisor)

        ttk.Separator(barra, orient="vertical").pack(side="left", fill="y", padx=10, pady=2)

        self._botao_icone(barra, "◀", "Página anterior", self.pagina_anterior)
        self.lbl_pagina = ttk.Label(barra, text="Página: -/-")
        self.lbl_pagina.pack(side="left", padx=6)
        self._botao_icone(barra, "▶", "Próxima página", self.proxima_pagina)
        self._botao_icone(barra, "🔍-", "Diminuir zoom", lambda: self.alterar_zoom(-0.2))
        self._botao_icone(barra, "🔍+", "Aumentar zoom", lambda: self.alterar_zoom(0.2))

        self._botao_icone(barra, "ℹ", "Sobre", self._abrir_sobre)
        if sv_ttk is not None:
            self._botao_icone(barra, "🌓", "Alternar tema claro/escuro", self._alternar_tema)

        ttk.Separator(barra, orient="vertical").pack(side="left", fill="y", padx=10, pady=2)
        self.btn_painel_lateral = self._botao_icone(
            barra, "☰", "Mostrar/ocultar painel de marcações", self._alternar_painel_lateral
        )

        ttk.Separator(barra, orient="vertical").pack(side="right", fill="y", padx=10, pady=2)
        self._botao_icone(barra, "📄", "Salvar PDF com comentários", self.exportar_pdf, side="right")
        self._botao_icone(barra, "📊", "Exportar Excel", self.exportar_excel, side="right")

        # Barra de ferramentas de marcação
        self._montar_barra_ferramentas()

        # Corpo: PanedWindow com canvas (esquerda) + painel lateral redimensionável/colapsável
        self.paned = ttk.PanedWindow(self.root, orient="horizontal")
        self.paned.pack(side="top", fill="both", expand=True)

        canvas_frame = ttk.Frame(self.paned)
        self.paned.add(canvas_frame, weight=4)

        self.canvas = tk.Canvas(
            canvas_frame, bg="#666666", cursor="crosshair",
            borderwidth=0, highlightthickness=1, highlightbackground=COR_BORDA_SUAVE,
        )
        hbar = ttk.Scrollbar(canvas_frame, orient="horizontal", command=self.canvas.xview)
        vbar = ttk.Scrollbar(canvas_frame, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(xscrollcommand=hbar.set, yscrollcommand=vbar.set)
        hbar.pack(side="bottom", fill="x")
        vbar.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True, padx=(0, 4))
        self.canvas.bind("<ButtonPress-1>", self._ao_pressionar)
        self.canvas.bind("<B1-Motion>", self._ao_arrastar)
        self.canvas.bind("<ButtonRelease-1>", self._ao_soltar)
        self.canvas.bind("<Double-Button-1>", self._ao_duplo_clique_canvas)

        self.lateral = ttk.Frame(self.paned, width=380)
        if self._lateral_visivel:
            self.paned.add(self.lateral, weight=1)

        cabecalho_lateral = ttk.Frame(self.lateral)
        cabecalho_lateral.pack(fill="x", padx=8, pady=(8, 4))
        ttk.Label(cabecalho_lateral, text="Marcações inseridas", font=("Segoe UI", 10, "bold")).pack(
            side="left"
        )
        self.btn_colapsar = ttk.Button(
            cabecalho_lateral, text="▶" if self._lateral_visivel else "◀",
            width=3, command=self._alternar_painel_lateral,
        )
        self.btn_colapsar.pack(side="right")
        Tooltip(self.btn_colapsar, "Colapsar/expandir painel lateral")

        filtro_frame = ttk.Frame(self.lateral)
        filtro_frame.pack(fill="x", padx=8, pady=(0, 4))
        ttk.Label(filtro_frame, text="Filtrar por status:").pack(side="left")
        self.filtro_status_var = tk.StringVar(value="Todos")
        ttk.Combobox(
            filtro_frame, textvariable=self.filtro_status_var,
            values=["Todos"] + STATUS_REVISAO, state="readonly", width=14,
        ).pack(side="left", padx=4)
        self.filtro_status_var.trace_add("write", lambda *a: self._atualizar_lista_comentarios())

        colunas = ("tipo", "pagina", "classificacao", "resumo", "status")
        self.tree = ttk.Treeview(self.lateral, columns=colunas, show="headings", height=23)
        self.tree.heading("tipo", text="Tipo")
        self.tree.heading("pagina", text="Pág.")
        self.tree.heading("classificacao", text="Classificação")
        self.tree.heading("resumo", text="Comentário")
        self.tree.heading("status", text="Status")
        self.tree.column("tipo", width=80, anchor="center")
        self.tree.column("pagina", width=32, anchor="center")
        self.tree.column("classificacao", width=100)
        self.tree.column("resumo", width=110)
        self.tree.column("status", width=90, anchor="center")
        self.tree.pack(fill="both", expand=True, padx=8)
        self.tree.bind("<Double-1>", self.editar_comentario_selecionado)
        self.tree.bind("<<TreeviewSelect>>", lambda e: self.renderizar_pagina())
        self.tree.bind("<Button-3>", self._menu_contexto_tree)
        self._configurar_tags_status()

        btns_lista = ttk.Frame(self.lateral)
        btns_lista.pack(fill="x", padx=8, pady=8)
        ttk.Button(btns_lista, text="Editar", command=self.editar_comentario_selecionado).pack(
            side="left", expand=True, fill="x", padx=2
        )
        ttk.Button(btns_lista, text="Excluir", command=self.excluir_comentario_selecionado).pack(
            side="left", expand=True, fill="x", padx=2
        )
        ttk.Button(btns_lista, text="Ir até", command=self.ir_ate_comentario).pack(
            side="left", expand=True, fill="x", padx=2
        )

        self.lbl_status = ttk.Label(self.root, text="Abra um PDF para começar.", anchor="w")
        self.lbl_status.pack(side="bottom", fill="x", padx=8, pady=4)

    def _botao_icone(self, parent, icone, tooltip_texto, comando, side="left"):
        btn = ttk.Button(parent, text=icone, width=4, command=comando)
        btn.pack(side=side, padx=3)
        Tooltip(btn, tooltip_texto)
        return btn

    def _montar_menu(self):
        menubar = tk.Menu(self.root)

        menu_arquivo = tk.Menu(menubar, tearoff=0)
        menu_arquivo.add_command(label="Abrir PDF", command=self.abrir_pdf)
        menu_arquivo.add_command(label="Salvar projeto", command=self.salvar_projeto)
        menu_arquivo.add_command(label="Abrir projeto", command=self.abrir_projeto)
        menu_arquivo.add_separator()
        menu_arquivo.add_command(label="Salvar PDF com comentários", command=self.exportar_pdf)
        menu_arquivo.add_command(label="Exportar Excel", command=self.exportar_excel)
        menu_arquivo.add_separator()
        menu_arquivo.add_command(label="Sair", command=self.root.quit)
        menubar.add_cascade(label="Arquivo", menu=menu_arquivo)

        menu_ajuda = tk.Menu(menubar, tearoff=0)
        menu_ajuda.add_command(label="Sobre", command=self._abrir_sobre)
        menubar.add_cascade(label="Ajuda", menu=menu_ajuda)

        self.root.config(menu=menubar)

    def _abrir_sobre(self):
        SobreDialog(self.root)

    def _alternar_painel_lateral(self):
        if self._lateral_visivel:
            self.paned.forget(self.lateral)
            self._lateral_visivel = False
            self.btn_colapsar.config(text="◀")
        else:
            self.paned.add(self.lateral, weight=1)
            self._lateral_visivel = True
            self.btn_colapsar.config(text="▶")

        if hasattr(self, "btn_painel_lateral"):
            self.btn_painel_lateral.state(["pressed"] if self._lateral_visivel else ["!pressed"])

    def _alternar_tema(self):
        if sv_ttk is None:
            return
        atual = sv_ttk.get_theme()
        sv_ttk.set_theme("dark" if atual == "light" else "light")

    def _montar_barra_ferramentas(self):
        barra2 = ttk.Frame(self.root, borderwidth=1, relief="groove")
        barra2.pack(side="top", fill="x", padx=6, pady=(0, 6))

        ttk.Label(barra2, text="Ferramenta:", font=("Segoe UI", 9, "bold")).pack(
            side="left", padx=(6, 8), pady=4
        )

        self.modo_var = tk.StringVar(value="comentario")
        for tool_id, info in FERRAMENTAS.items():
            rb = ttk.Radiobutton(
                barra2,
                text=info["icone"],
                value=tool_id,
                variable=self.modo_var,
                style="Toolbutton",
                width=3,
                command=self._atualizar_opcoes_ferramenta,
            )
            rb.pack(side="left", padx=1, pady=4)
            Tooltip(rb, info["label"])

        ttk.Separator(barra2, orient="vertical").pack(side="left", fill="y", padx=10, pady=2)

        self.frame_opcoes = ttk.Frame(barra2)
        self.frame_opcoes.pack(side="left", padx=4, pady=4)

        # estado atual do estilo (usado como padrão ao criar a PRÓXIMA marcação;
        # também pode ser reaplicado a um elemento já selecionado)
        self.var_fontsize = tk.IntVar(value=12)
        self.cor_texto_livre = "#000000"
        self.var_cor_realce = tk.StringVar(value="Amarelo")

        self.estilo_cor_borda = "#FF0000"
        self.estilo_cor_preenchimento = None
        self.estilo_espessura = tk.IntVar(value=2)
        self.estilo_tracejado = tk.StringVar(value="Contínua")
        self.estilo_opacidade_borda = tk.IntVar(value=100)
        self.estilo_opacidade_preenchimento = tk.IntVar(value=100)
        self._var_seta_linha = tk.BooleanVar(value=False)

        self.caneta_cor = "#FF0000"
        self.caneta_espessura = tk.IntVar(value=3)
        self.caneta_opacidade = tk.IntVar(value=100)

        self.carimbo_imagem_path = None

        self._atualizar_opcoes_ferramenta()

    def _atualizar_opcoes_ferramenta(self):
        if hasattr(self, "canvas"):
            tool_atual = self.modo_var.get()
            cursor = "fleur" if tool_atual == "mao" else (
                "arrow" if tool_atual == "selecionar" else "crosshair"
            )
            self.canvas.config(cursor=cursor)

        for w in self.frame_opcoes.winfo_children():
            w.destroy()

        tool = self.modo_var.get()

        if tool == "texto_livre":
            ttk.Label(self.frame_opcoes, text="Tamanho:").pack(side="left")
            ttk.Spinbox(
                self.frame_opcoes, from_=6, to=72, width=4, textvariable=self.var_fontsize
            ).pack(side="left", padx=(2, 10))
            self._campo_cor_ferramenta(self.frame_opcoes, "Cor do texto:", "cor_texto_livre")

        elif tool == "realce":
            ttk.Label(self.frame_opcoes, text="Cor do realce:").pack(side="left")
            ttk.Combobox(
                self.frame_opcoes,
                textvariable=self.var_cor_realce,
                values=list(CORES_REALCE.keys()),
                state="readonly",
                width=10,
            ).pack(side="left", padx=4)

        elif tool in ("sublinhado", "tachado"):
            self._campo_cor_ferramenta(self.frame_opcoes, "Cor:", "estilo_cor_borda")

        elif tool == "caneta":
            self._campo_cor_ferramenta(self.frame_opcoes, "Cor:", "caneta_cor")
            ttk.Label(self.frame_opcoes, text="Espessura:").pack(side="left", padx=(10, 0))
            ttk.Spinbox(
                self.frame_opcoes, from_=1, to=12, width=3, textvariable=self.caneta_espessura
            ).pack(side="left", padx=(2, 10))
            ttk.Label(self.frame_opcoes, text="Opacidade:").pack(side="left")
            ttk.Scale(
                self.frame_opcoes, from_=0, to=100, orient="horizontal",
                variable=self.caneta_opacidade, length=100,
            ).pack(side="left", padx=4)

        elif tool in FERRAMENTAS_COM_BORDA:
            self._campo_cor_ferramenta(self.frame_opcoes, "Cor da borda:", "estilo_cor_borda")
            if tool in FERRAMENTAS_COM_PREENCHIMENTO:
                self._campo_preenchimento(self.frame_opcoes)
            ttk.Label(self.frame_opcoes, text="Espessura:").pack(side="left", padx=(10, 0))
            ttk.Spinbox(
                self.frame_opcoes, from_=1, to=10, width=3, textvariable=self.estilo_espessura
            ).pack(side="left", padx=(2, 10))
            ttk.Label(self.frame_opcoes, text="Traço:").pack(side="left")
            ttk.Combobox(
                self.frame_opcoes,
                textvariable=self.estilo_tracejado,
                values=ESTILOS_TRACEJADO,
                state="readonly",
                width=10,
            ).pack(side="left", padx=4)
            if tool == "linha":
                ttk.Checkbutton(
                    self.frame_opcoes, text="Com seta na ponta", variable=self._var_seta_linha
                ).pack(side="left", padx=(10, 0))
            if tool in FERRAMENTAS_COM_PREENCHIMENTO:
                ttk.Label(self.frame_opcoes, text="Opac. borda:").pack(side="left", padx=(10, 0))
                ttk.Scale(
                    self.frame_opcoes, from_=0, to=100, orient="horizontal",
                    variable=self.estilo_opacidade_borda, length=70,
                ).pack(side="left", padx=2)
                ttk.Label(self.frame_opcoes, text="Opac. preench.:").pack(side="left", padx=(6, 0))
                ttk.Scale(
                    self.frame_opcoes, from_=0, to=100, orient="horizontal",
                    variable=self.estilo_opacidade_preenchimento, length=70,
                ).pack(side="left", padx=2)

        elif tool == "carimbo":
            ttk.Button(
                self.frame_opcoes, text="Configurar imagem do carimbo...",
                command=self._configurar_carimbo,
            ).pack(side="left")
            nome_img = self._descricao_imagem_carimbo_atual()
            ttk.Label(self.frame_opcoes, text=f"Imagem atual: {nome_img}").pack(
                side="left", padx=8
            )

        elif tool == "selecionar":
            ttk.Label(
                self.frame_opcoes,
                text="Clique seleciona; Shift/Ctrl+clique adiciona; arraste em área vazia "
                     "seleciona por retângulo. Duplo clique edita, Delete exclui.",
            ).pack(side="left")
            ttk.Button(
                self.frame_opcoes, text="Aplicar estilo ao selecionado",
                command=self._aplicar_estilo_ao_selecionado,
            ).pack(side="left", padx=10)

        elif tool == "mao":
            ttk.Label(
                self.frame_opcoes, text="Clique e arraste para rolar o documento."
            ).pack(side="left")

    def _descricao_imagem_carimbo_atual(self):
        if self.carimbo_imagem_path:
            return os.path.basename(self.carimbo_imagem_path)
        if os.path.exists(CAMINHO_LOGO_PADRAO):
            return f"{NOME_ARQUIVO_LOGO_PADRAO} (logo PREVENT, automática)"
        return "(padrão gerado por código)"

    def _campo_cor_ferramenta(self, parent, rotulo, attr_nome):
        ttk.Label(parent, text=rotulo).pack(side="left")
        swatch = tk.Label(
            parent, text="    ", bg=getattr(self, attr_nome), relief="solid", borderwidth=1
        )
        swatch.pack(side="left", padx=(4, 4))

        def escolher():
            cor = colorchooser.askcolor(color=getattr(self, attr_nome), title=rotulo)
            if cor and cor[1]:
                setattr(self, attr_nome, cor[1])
                swatch.config(bg=cor[1])

        ttk.Button(parent, text="Escolher...", command=escolher).pack(side="left")

    def _campo_preenchimento(self, parent):
        ttk.Label(parent, text="Preenchimento:").pack(side="left", padx=(10, 0))
        cor_atual = self.estilo_cor_preenchimento
        swatch = tk.Label(
            parent, text="    ", relief="solid", borderwidth=1, bg=cor_atual or "#F0F0F0"
        )
        swatch.pack(side="left", padx=(4, 4))
        sem_preench_var = tk.BooleanVar(value=self.estilo_cor_preenchimento is None)

        def alternar():
            if sem_preench_var.get():
                self.estilo_cor_preenchimento = None
                swatch.config(bg="#F0F0F0")
            else:
                self.estilo_cor_preenchimento = self.estilo_cor_preenchimento or "#FFFFFF"
                swatch.config(bg=self.estilo_cor_preenchimento)

        def escolher():
            cor = colorchooser.askcolor(
                color=self.estilo_cor_preenchimento or "#FFFFFF", title="Cor de preenchimento"
            )
            if cor and cor[1]:
                self.estilo_cor_preenchimento = cor[1]
                sem_preench_var.set(False)
                swatch.config(bg=cor[1])

        ttk.Checkbutton(
            parent, text="Sem preenchimento", variable=sem_preench_var, command=alternar
        ).pack(side="left")
        ttk.Button(parent, text="Escolher...", command=escolher).pack(side="left", padx=(4, 0))

    def _atalhos(self):
        self.root.bind("<Right>", lambda e: self.proxima_pagina())
        self.root.bind("<Left>", lambda e: self.pagina_anterior())
        self.root.bind("<Delete>", self._excluir_via_teclado)
        self.root.bind("<BackSpace>", self._excluir_via_teclado)
        self.root.bind("<Escape>", self._tecla_escape)
        for i in range(4):
            self.root.bind(f"<Key-{i + 1}>", lambda e, idx=i: self._status_via_teclado(idx))

    def _tecla_escape(self, event=None):
        self._arrasto_inicio = None
        self._marquee_inicio = None
        self._mover_idx = None
        self._mover_original_grupo = {}
        self._pontos_caneta_canvas = []
        self._limpar_preview()
        self.tree.selection_remove(*self.tree.selection())
        self.modo_var.set("selecionar")
        self._atualizar_opcoes_ferramenta()
        self.renderizar_pagina()

    # ------------------------- revisão colaborativa (status) ------------------------- #
    def _configurar_tags_status(self):
        self.tree.tag_configure("status_pendente", background="#FFFFFF")
        self.tree.tag_configure("status_corrigido", background="#DFF5DF")
        self.tree.tag_configure("status_nao_corrigido", background="#FCE0E0")
        self.tree.tag_configure("status_reincidente", background="#FFE7C2")
        self.tree.tag_configure("status_correto", background="#FFF9C4")

    def _identificar_revisor(self):
        dialog = RevisorDialog(self.root, self.revisor_atual or "")
        if dialog.result:
            self.revisor_atual = dialog.result
            self.lbl_status.config(text=f"Revisor identificado: {self.revisor_atual}")

    def _garantir_revisor(self):
        if self.revisor_atual:
            return self.revisor_atual
        dialog = RevisorDialog(self.root)
        if dialog.result:
            self.revisor_atual = dialog.result
        return self.revisor_atual

    def _aplicar_status(self, item, novo_status):
        item["status"] = novo_status
        if novo_status != "Pendente":
            revisor = self._garantir_revisor()
            item["revisado_por"] = revisor or item.get("revisado_por", "")
            item["data_revisao"] = hoje_str()
        else:
            item.setdefault("revisado_por", "")
            item.setdefault("data_revisao", "")

    def _aplicar_verificacao_correcao(self, item, novo_valor):
        item["verificacao_correcao"] = bool(novo_valor) and item.get("status") == "Corrigido"
        if item["verificacao_correcao"]:
            revisor = self._garantir_revisor()
            item["verificado_correcao_por"] = revisor or item.get("verificado_correcao_por", "")
            item["data_verificacao_correcao"] = hoje_str()
        else:
            item.setdefault("verificado_correcao_por", "")
            item.setdefault("data_verificacao_correcao", "")

    def _marcar_como_correto(self, item):
        item["status"] = "Correto"
        item["classificacao"] = ""
        item["texto"] = TEXTO_ITEM_CORRETO
        item["verificacao_correcao"] = False
        item["verificado_correcao_por"] = ""
        item["data_verificacao_correcao"] = ""
        revisor = self._garantir_revisor()
        item["revisado_por"] = revisor or ""
        item["data_revisao"] = hoje_str()

    def _cor_ativa_da_ferramenta(self, tool):
        """Cor de traço/preenchimento atualmente selecionada no painel de estilo
        para a ferramenta em questão (ou None se a ferramenta não tiver cor)."""
        if tool == "texto_livre":
            return self.cor_texto_livre
        if tool == "realce":
            return CORES_REALCE[self.var_cor_realce.get()]
        if tool in ("sublinhado", "tachado"):
            return self.estilo_cor_borda
        if tool == "caneta":
            return self.caneta_cor
        if tool in FERRAMENTAS_COM_BORDA:
            if self.estilo_cor_preenchimento and cor_e_amarela(self.estilo_cor_preenchimento):
                return self.estilo_cor_preenchimento
            return self.estilo_cor_borda
        return None

    def _alternar_verificacao_selecionado(self):
        idx = self._indice_selecionado()
        if idx is None:
            return
        c = self.comentarios[idx]
        if c.get("status") != "Corrigido":
            messagebox.showinfo(
                "Aviso",
                "Só é possível marcar a verificação da correção quando o "
                "status do item for 'Corrigido'.",
            )
            return
        self._aplicar_verificacao_correcao(c, not c.get("verificacao_correcao", False))
        self._atualizar_lista_comentarios()
        self.renderizar_pagina()

    def _definir_status_selecionado(self, status):
        indices = self._indices_selecionados()
        if not indices:
            return
        for idx in indices:
            self._aplicar_status(self.comentarios[idx], status)
        self._atualizar_lista_comentarios()
        self.renderizar_pagina()

    def _status_via_teclado(self, indice):
        foco = self.root.focus_get()
        if isinstance(foco, (tk.Entry, ttk.Entry, tk.Spinbox, ttk.Spinbox, tk.Text)):
            return
        if not self._indices_selecionados():
            return
        self._definir_status_selecionado(STATUS_REVISAO[indice])

    def _menu_contexto_tree(self, event):
        iid = self.tree.identify_row(event.y)
        if not iid:
            return
        if iid not in self.tree.selection():
            self.tree.selection_set(iid)
        self.renderizar_pagina()

        menu = tk.Menu(self.root, tearoff=0)
        for status in STATUS_REVISAO:
            menu.add_command(
                label=f"Status: {status}",
                command=lambda s=status: self._definir_status_selecionado(s),
            )

        c = self.comentarios[int(iid)]
        if c.get("tipo") != "carimbo":
            menu.add_separator()
            marcado = c.get("verificacao_correcao", False)
            rotulo_verificacao = ("Desmarcar" if marcado else "Marcar") + " verificação da correção"
            menu.add_command(label=rotulo_verificacao, command=self._alternar_verificacao_selecionado)

        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    # ------------------------- abrir / navegar PDF ------------------------- #
    def abrir_pdf(self):
        caminho = filedialog.askopenfilename(
            title="Selecione o PDF", filetypes=[("Arquivos PDF", "*.pdf")]
        )
        if not caminho:
            return
        self.doc = fitz.open(caminho)
        self.pdf_path = caminho
        self.pagina_atual = 0
        self.comentarios = []
        self._atualizar_lista_comentarios()

        nome_doc = os.path.splitext(os.path.basename(caminho))[0]
        dialog = DadosDocumentoDialog(
            self.root, {"codigo_doc": nome_doc, "data_elaboracao": hoje_str()}
        )
        if dialog.result:
            self.dados_documento = dialog.result

        self.renderizar_pagina()
        self.lbl_status.config(text=f"PDF aberto: {caminho}")

    def editar_dados_documento(self):
        dialog = DadosDocumentoDialog(self.root, self.dados_documento)
        if dialog.result:
            self.dados_documento = dialog.result

    def renderizar_pagina(self):
        if not self.doc:
            return
        pagina = self.doc[self.pagina_atual]
        mat = fitz.Matrix(self.zoom, self.zoom)
        pix = pagina.get_pixmap(matrix=mat)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        self.tk_img = ImageTk.PhotoImage(img)

        self.canvas.delete("all")
        self.canvas.create_image(0, 0, anchor="nw", image=self.tk_img)
        self.canvas.config(scrollregion=(0, 0, pix.width, pix.height))

        # desenha marcadores das marcações já existentes nesta página
        for idx, c in enumerate(self.comentarios):
            if c["pagina"] == self.pagina_atual:
                self._desenhar_marcador(c, idx + 1)
                self._desenhar_confirmacao(c)

        for idx_sel in self._indices_selecionados():
            if idx_sel < len(self.comentarios):
                c_sel = self.comentarios[idx_sel]
                if c_sel["pagina"] == self.pagina_atual:
                    self._desenhar_selecao(c_sel)

        total_pag = len(self.doc)
        self.lbl_pagina.config(text=f"Página: {self.pagina_atual + 1}/{total_pag}")

    def _ponto_canvas(self, x, y):
        return x * self.zoom, y * self.zoom

    def _desenhar_selecao(self, c):
        x0, y0, x1, y1 = bbox_do_elemento(c)
        cx0, cy0 = self._ponto_canvas(x0, y0)
        cx1, cy1 = self._ponto_canvas(x1, y1)
        margem = 6
        cx0, cy0, cx1, cy1 = cx0 - margem, cy0 - margem, cx1 + margem, cy1 + margem
        self.canvas.create_rectangle(cx0, cy0, cx1, cy1, outline="#0078D4", width=2, dash=(6, 3))
        for hx, hy in ((cx0, cy0), (cx1, cy0), (cx0, cy1), (cx1, cy1)):
            self.canvas.create_rectangle(hx - 3, hy - 3, hx + 3, hy + 3, fill="#0078D4", outline="")

    def _desenhar_marcador(self, c, numero):
        tipo = c.get("tipo", "comentario")
        cor_selo = cor_status(c)

        if tipo == "comentario":
            cx, cy = self._ponto_canvas(c["x"], c["y"])
            self.canvas.create_oval(
                cx - 9, cy - 9, cx + 9, cy + 9, fill=cor_selo, outline="black", width=1
            )
            self.canvas.create_text(cx, cy, text=str(numero), font=("Segoe UI", 8, "bold"))
            return

        if tipo == "texto_livre":
            x, y = self._ponto_canvas(c["x"], c["y"])
            fontsize = c.get("fontsize", 12)
            cor = c.get("cor", "#000000")
            self.canvas.create_text(
                x, y, text=c["texto"], anchor="nw", fill=cor,
                font=("Segoe UI", max(6, int(fontsize * self.zoom))), width=400,
            )
            self._desenhar_selo(x, y, numero, cor_selo)
            return

        if tipo == "caneta":
            pontos_canvas = []
            for px, py in c["pontos"]:
                cx, cy = self._ponto_canvas(px, py)
                pontos_canvas.extend([cx, cy])
            if len(pontos_canvas) >= 4:
                cor_traco = hex_com_opacidade_simulada(
                    c.get("cor", "#FF0000"), c.get("opacidade", 100)
                )
                self.canvas.create_line(
                    *pontos_canvas, fill=cor_traco, width=c.get("espessura", 3),
                    smooth=True, capstyle="round", joinstyle="round",
                )
            self._desenhar_selo(pontos_canvas[0], pontos_canvas[1], numero, cor_selo)
            return

        if tipo == "carimbo":
            x0, y0 = self._ponto_canvas(min(c["x0"], c["x1"]), min(c["y0"], c["y1"]))
            x1, y1 = self._ponto_canvas(max(c["x0"], c["x1"]), max(c["y0"], c["y1"]))
            self.canvas.create_rectangle(x0, y0, x1, y1, outline="#00529B", width=2, dash=(3, 2))
            self.canvas.create_text(
                (x0 + x1) / 2, (y0 + y1) / 2, text="CARIMBO",
                fill="#00529B", font=("Segoe UI", 9, "bold")
            )
            self._desenhar_selo(x0, y0, numero, cor_selo)
            return

        if tipo in ("linha", "seta"):
            x0, y0 = self._ponto_canvas(c["x0"], c["y0"])
            x1, y1 = self._ponto_canvas(c["x1"], c["y1"])
            cor = c.get("cor_borda", "#ff0000")
            largura = espessura_efetiva(c)
            dash = DASHES_CANVAS.get(c.get("tracejado", "Contínua"))
            seta = tk.LAST if (tipo == "seta" or c.get("seta")) else tk.NONE
            self.canvas.create_line(x0, y0, x1, y1, fill=cor, width=largura, dash=dash, arrow=seta)
            self._desenhar_selo(x0, y0, numero, cor_selo)
            return

        x0, y0 = self._ponto_canvas(min(c["x0"], c["x1"]), min(c["y0"], c["y1"]))
        x1, y1 = self._ponto_canvas(max(c["x0"], c["x1"]), max(c["y0"], c["y1"]))

        if tipo == "retangulo":
            cor = hex_com_opacidade_simulada(
                c.get("cor_borda", "#FF0000"), c.get("opacidade_borda", 100)
            )
            largura = espessura_efetiva(c)
            dash = DASHES_CANVAS.get(c.get("tracejado", "Contínua")) or (4, 2)
            preench = c.get("cor_preenchimento")
            preench_final = (
                hex_com_opacidade_simulada(preench, c.get("opacidade_preenchimento", 100))
                if preench else ""
            )
            self.canvas.create_rectangle(
                x0, y0, x1, y1, outline=cor, width=largura, dash=dash, fill=preench_final
            )
        elif tipo == "elipse":
            cor = hex_com_opacidade_simulada(
                c.get("cor_borda", "#000000"), c.get("opacidade_borda", 100)
            )
            largura = espessura_efetiva(c)
            dash = DASHES_CANVAS.get(c.get("tracejado", "Contínua")) or (4, 2)
            preench = c.get("cor_preenchimento")
            preench_final = (
                hex_com_opacidade_simulada(preench, c.get("opacidade_preenchimento", 100))
                if preench else ""
            )
            self.canvas.create_oval(
                x0, y0, x1, y1, outline=cor, width=largura, dash=dash, fill=preench_final
            )
        elif tipo == "nuvem":
            cor = hex_com_opacidade_simulada(
                c.get("cor_borda", "#FF0000"), c.get("opacidade_borda", 100)
            )
            largura = espessura_efetiva(c)
            self._desenhar_bolhas_nuvem(x0, y0, x1, y1, cor, largura)
        elif tipo == "sublinhado":
            self.canvas.create_line(x0, y1, x1, y1, fill=c.get("cor", "#0000ff"), width=2)
        elif tipo == "tachado":
            ym = (y0 + y1) / 2
            self.canvas.create_line(x0, ym, x1, ym, fill=c.get("cor", "#ff0000"), width=2)
        elif tipo == "realce":
            cor = c.get("cor", "#FFFF00")
            self.canvas.create_rectangle(x0, y0, x1, y1, fill=cor, outline="", stipple="gray50")

        self._desenhar_selo(x0, y0, numero, cor_selo)

    def _desenhar_bolhas_nuvem(self, x0, y0, x1, y1, cor, largura):
        """Aproximação visual de nuvem de revisão no canvas (bolhas ao longo do
        perímetro); a exportação final usa o efeito 'Cloudy' real do PyMuPDF."""
        raio = 10

        def bolhas_no_segmento(ax, ay, bx, by):
            comprimento = ((bx - ax) ** 2 + (by - ay) ** 2) ** 0.5
            n = max(1, int(comprimento // (raio * 1.4)))
            for i in range(n):
                t0, t1 = i / n, (i + 1) / n
                px0, py0 = ax + (bx - ax) * t0, ay + (by - ay) * t0
                px1, py1 = ax + (bx - ax) * t1, ay + (by - ay) * t1
                mx, my = (px0 + px1) / 2, (py0 + py1) / 2
                self.canvas.create_arc(
                    mx - raio, my - raio, mx + raio, my + raio,
                    start=0, extent=180, style="arc", outline=cor, width=largura,
                )

        bolhas_no_segmento(x0, y0, x1, y0)
        bolhas_no_segmento(x1, y0, x1, y1)
        bolhas_no_segmento(x1, y1, x0, y1)
        bolhas_no_segmento(x0, y1, x0, y0)

    def _desenhar_selo(self, x, y, numero, cor="#ffcc00"):
        self.canvas.create_oval(x - 8, y - 8, x + 8, y + 8, fill=cor, outline="black")
        self.canvas.create_text(x, y, text=str(numero), font=("Segoe UI", 8, "bold"))

    def _desenhar_confirmacao(self, c):
        """Elipse azul (confirmação de execução da correção, status=='Corrigido') com
        X verde dentro (validação/verificação da correção), padrão PP-ENG-002 5.1.4.3.
        Não se aplica ao Carimbo. Puramente derivado do estado atual: não é um desenho
        permanente, então desaparece sozinho se o status voltar a mudar."""
        if c.get("tipo") == "carimbo" or c.get("status") != "Corrigido":
            return
        x0, y0, x1, y1 = bbox_do_elemento(c)
        cx0, cy0 = self._ponto_canvas(x0, y0)
        cx1, cy1 = self._ponto_canvas(x1, y1)
        m = MARGEM_CONFIRMACAO
        cx0, cy0, cx1, cy1 = cx0 - m, cy0 - m, cx1 + m, cy1 + m
        self.canvas.create_oval(cx0, cy0, cx1, cy1, outline=COR_AZUL_CONFIRMACAO, width=2)
        if c.get("verificacao_correcao"):
            self.canvas.create_line(cx0, cy0, cx1, cy1, fill=COR_VERDE_VALIDACAO, width=2)
            self.canvas.create_line(cx1, cy0, cx0, cy1, fill=COR_VERDE_VALIDACAO, width=2)

    def proxima_pagina(self):
        if self.doc and self.pagina_atual < len(self.doc) - 1:
            self.pagina_atual += 1
            self.renderizar_pagina()

    def pagina_anterior(self):
        if self.doc and self.pagina_atual > 0:
            self.pagina_atual -= 1
            self.renderizar_pagina()

    def alterar_zoom(self, delta):
        if not self.doc:
            return
        novo = round(self.zoom + delta, 2)
        if 0.4 <= novo <= 4.0:
            self.zoom = novo
            self.renderizar_pagina()

    # ------------------------- interação com o canvas ------------------------- #
    def _ao_pressionar(self, event):
        if not self.doc:
            messagebox.showinfo("Aviso", "Abra um PDF primeiro.")
            return
        tool = self.modo_var.get()
        interacao = FERRAMENTAS[tool]["interacao"]

        if interacao == "mao":
            self.canvas.scan_mark(event.x, event.y)
            return
        if interacao == "selecionar":
            self._selecionar_ao_pressionar(event)
            return
        if interacao == "caneta":
            cx, cy = self.canvas.canvasx(event.x), self.canvas.canvasy(event.y)
            self._pontos_caneta_canvas = [(cx, cy)]
            return
        if interacao == "clique":
            if tool == "carimbo":
                self._criar_carimbo(event)
            else:
                self._criar_marcacao_clique(tool, event)
            return
        self._arrasto_inicio = (self.canvas.canvasx(event.x), self.canvas.canvasy(event.y))

    def _ao_arrastar(self, event):
        tool = self.modo_var.get()
        interacao = FERRAMENTAS[tool]["interacao"]

        if interacao == "mao":
            self.canvas.scan_dragto(event.x, event.y, gain=1)
            return
        if interacao == "selecionar":
            self._selecionar_ao_arrastar(event)
            return
        if interacao == "caneta":
            if not self._pontos_caneta_canvas:
                return
            cx, cy = self.canvas.canvasx(event.x), self.canvas.canvasy(event.y)
            self._pontos_caneta_canvas.append((cx, cy))
            self._redesenhar_preview_caneta()
            return
        if self._arrasto_inicio is None:
            return
        cx, cy = self.canvas.canvasx(event.x), self.canvas.canvasy(event.y)
        self._desenhar_preview(tool, self._arrasto_inicio, (cx, cy))

    def _ao_soltar(self, event):
        tool = self.modo_var.get()
        interacao = FERRAMENTAS[tool]["interacao"]

        if interacao == "mao":
            return
        if interacao == "selecionar":
            if self._marquee_inicio is not None:
                cx0, cy0 = self._marquee_inicio
                cx1, cy1 = self.canvas.canvasx(event.x), self.canvas.canvasy(event.y)
                self._limpar_preview()
                self._finalizar_marquee(cx0, cy0, cx1, cy1)
                self._marquee_inicio = None
            self._mover_idx = None
            self._mover_original_grupo = {}
            return
        if interacao == "caneta":
            self._finalizar_caneta()
            return
        if self._arrasto_inicio is None:
            return
        cx0, cy0 = self._arrasto_inicio
        cx1, cy1 = self.canvas.canvasx(event.x), self.canvas.canvasy(event.y)
        self._arrasto_inicio = None
        self._limpar_preview()
        if abs(cx1 - cx0) < 3 and abs(cy1 - cy0) < 3:
            return  # arrasto acidental, ignora
        self._criar_marcacao_arrasto(tool, cx0, cy0, cx1, cy1)

    def _ao_duplo_clique_canvas(self, event):
        if self.modo_var.get() != "selecionar":
            return
        if self._indice_selecionado() is not None:
            self.editar_comentario_selecionado()

    def _excluir_via_teclado(self, event=None):
        if self.modo_var.get() != "selecionar":
            return
        if not self._indices_selecionados():
            return
        self.excluir_comentario_selecionado()

    # ---- ferramenta Caneta (desenho livre) ---- #
    def _redesenhar_preview_caneta(self):
        if self._preview_caneta_id is not None:
            self.canvas.delete(self._preview_caneta_id)
            self._preview_caneta_id = None
        if len(self._pontos_caneta_canvas) < 2:
            return
        achatado = [coord for ponto in self._pontos_caneta_canvas for coord in ponto]
        cor = hex_com_opacidade_simulada(self.caneta_cor, self.caneta_opacidade.get())
        self._preview_caneta_id = self.canvas.create_line(
            *achatado, fill=cor, width=self.caneta_espessura.get(),
            smooth=True, capstyle="round", joinstyle="round",
        )

    def _finalizar_caneta(self):
        pontos_canvas = self._pontos_caneta_canvas
        self._pontos_caneta_canvas = []
        if self._preview_caneta_id is not None:
            self.canvas.delete(self._preview_caneta_id)
            self._preview_caneta_id = None
        if len(pontos_canvas) < 2:
            return

        pontos_pdf = [[cx / self.zoom, cy / self.zoom] for cx, cy in pontos_canvas]

        base = {
            "tipo": "caneta",
            "pagina": self.pagina_atual,
            "pontos": pontos_pdf,
            "cor": self.caneta_cor,
            "espessura": self.caneta_espessura.get(),
            "opacidade": self.caneta_opacidade.get(),
            "data": hoje_str(),
        }

        if cor_e_amarela(self.caneta_cor):
            self._marcar_como_correto(base)
            self.comentarios.append(base)
            self._atualizar_lista_comentarios()
            self.renderizar_pagina()
            return

        dialog = ComentarioDialog(self.root, self.pagina_atual, tipo="caneta")
        if not dialog.result:
            return  # marcação descartada se o usuário cancelar

        base["texto"] = dialog.result["texto"]
        base["classificacao"] = dialog.result["classificacao"]
        self._aplicar_status(base, dialog.result.get("status", "Pendente"))
        self._aplicar_verificacao_correcao(base, dialog.result.get("verificacao_correcao", False))

        self.comentarios.append(base)
        self._atualizar_lista_comentarios()
        self.renderizar_pagina()

    # ---- modo Selecionar: hit-test, seleção múltipla, marquee e movimentação ---- #
    def _elemento_no_ponto(self, px, py):
        margem = 6
        for idx in range(len(self.comentarios) - 1, -1, -1):
            c = self.comentarios[idx]
            if c["pagina"] != self.pagina_atual:
                continue
            x0, y0, x1, y1 = bbox_do_elemento(c)
            if x0 - margem <= px <= x1 + margem and y0 - margem <= py <= y1 + margem:
                return idx
        return None

    def _selecionar_ao_pressionar(self, event):
        cx, cy = self.canvas.canvasx(event.x), self.canvas.canvasy(event.y)
        px, py = cx / self.zoom, cy / self.zoom
        shift_ou_ctrl = bool(event.state & 0x0005)  # Shift (0x1) ou Control (0x4)

        idx_alvo = self._elemento_no_ponto(px, py)

        if idx_alvo is None:
            if not shift_ou_ctrl:
                self.tree.selection_remove(*self.tree.selection())
            self._marquee_inicio = (cx, cy)
            self._mover_idx = None
            self.renderizar_pagina()
            return

        self._marquee_inicio = None
        selecionados_atuais = set(self.tree.selection())
        iid_alvo = str(idx_alvo)
        if shift_ou_ctrl:
            if iid_alvo in selecionados_atuais:
                self.tree.selection_remove(iid_alvo)
            else:
                self.tree.selection_add(iid_alvo)
        elif iid_alvo not in selecionados_atuais:
            self.tree.selection_set(iid_alvo)
        # se já fazia parte de uma seleção múltipla, um clique simples nele
        # mantém o grupo (permite iniciar o arrasto do grupo inteiro)

        self._mover_idx = idx_alvo
        self._mover_inicio_px = (px, py)
        self._mover_original_grupo = {
            i: dict(self.comentarios[i]) for i in self._indices_selecionados()
        }
        self.renderizar_pagina()

    def _selecionar_ao_arrastar(self, event):
        cx, cy = self.canvas.canvasx(event.x), self.canvas.canvasy(event.y)
        if self._marquee_inicio is not None:
            self._desenhar_marquee(self._marquee_inicio, (cx, cy))
            return
        if self._mover_idx is None:
            return
        px, py = cx / self.zoom, cy / self.zoom
        dx = px - self._mover_inicio_px[0]
        dy = py - self._mover_inicio_px[1]

        for i, original in self._mover_original_grupo.items():
            c = self.comentarios[i]
            if "x" in original:
                c["x"] = original["x"] + dx
                c["y"] = original["y"] + dy
            if "x0" in original:
                c["x0"] = original["x0"] + dx
                c["y0"] = original["y0"] + dy
                c["x1"] = original["x1"] + dx
                c["y1"] = original["y1"] + dy
            if "pontos" in original:
                c["pontos"] = [[px0 + dx, py0 + dy] for px0, py0 in original["pontos"]]

        self.renderizar_pagina()

    def _desenhar_marquee(self, inicio, fim):
        self._limpar_preview()
        x0, y0 = inicio
        x1, y1 = fim
        item = self.canvas.create_rectangle(x0, y0, x1, y1, outline="#0078D4", dash=(4, 2))
        self._preview_ids.append(item)

    def _finalizar_marquee(self, cx0, cy0, cx1, cy1):
        if abs(cx1 - cx0) < 3 and abs(cy1 - cy0) < 3:
            return  # clique simples em área vazia, já tratado no _ao_pressionar
        px0, py0 = min(cx0, cx1) / self.zoom, min(cy0, cy1) / self.zoom
        px1, py1 = max(cx0, cx1) / self.zoom, max(cy0, cy1) / self.zoom
        novos_selecionados = []
        for idx, c in enumerate(self.comentarios):
            if c["pagina"] != self.pagina_atual:
                continue
            x0, y0, x1, y1 = bbox_do_elemento(c)
            if x0 <= px1 and x1 >= px0 and y0 <= py1 and y1 >= py0:
                novos_selecionados.append(str(idx))
        self.tree.selection_set(novos_selecionados)
        self.renderizar_pagina()

    def _aplicar_estilo_ao_selecionado(self):
        indices = self._indices_selecionados()
        if not indices:
            messagebox.showinfo("Aviso", "Selecione um elemento na lista ou no PDF primeiro.")
            return

        for idx in indices:
            c = self.comentarios[idx]
            tipo = c.get("tipo", "comentario")

            if tipo in FERRAMENTAS_COM_BORDA:
                c["cor_borda"] = self.estilo_cor_borda
                c["espessura"] = self.estilo_espessura.get()
                c["tracejado"] = self.estilo_tracejado.get()
            if tipo in FERRAMENTAS_COM_PREENCHIMENTO:
                c["cor_preenchimento"] = self.estilo_cor_preenchimento
                c["opacidade_borda"] = self.estilo_opacidade_borda.get()
                c["opacidade_preenchimento"] = self.estilo_opacidade_preenchimento.get()
            if tipo in ("sublinhado", "tachado"):
                c["cor"] = self.estilo_cor_borda
            if tipo == "realce":
                c["cor"] = CORES_REALCE[self.var_cor_realce.get()]
            if tipo == "texto_livre":
                c["cor"] = self.cor_texto_livre
                c["fontsize"] = self.var_fontsize.get()
            if tipo == "caneta":
                c["cor"] = self.caneta_cor
                c["espessura"] = self.caneta_espessura.get()
                c["opacidade"] = self.caneta_opacidade.get()

        self.renderizar_pagina()
        self.lbl_status.config(text=f"Estilo aplicado a {len(indices)} elemento(s).")

    def _desenhar_preview(self, tool, inicio, fim):
        self._limpar_preview()
        x0, y0 = inicio
        x1, y1 = fim

        if tool in FERRAMENTAS_COM_BORDA:
            cor = self.estilo_cor_borda
            largura = self.estilo_espessura.get()
            dash_canvas = DASHES_CANVAS.get(self.estilo_tracejado.get())
        else:
            cor, largura, dash_canvas = "#ff0000", 2, (4, 2)

        if tool in ("linha", "seta"):
            seta = tk.LAST if (tool == "seta" or self.var_seta_ativa()) else tk.NONE
            item = self.canvas.create_line(
                x0, y0, x1, y1, fill=cor, width=largura, dash=dash_canvas, arrow=seta
            )
            self._preview_ids.append(item)
        elif tool == "elipse":
            item = self.canvas.create_oval(
                x0, y0, x1, y1, outline=cor, width=largura, dash=dash_canvas or (4, 2)
            )
            self._preview_ids.append(item)
        elif tool == "nuvem":
            raio = 10

            def bolhas(ax, ay, bx, by):
                comprimento = ((bx - ax) ** 2 + (by - ay) ** 2) ** 0.5
                n = max(1, int(comprimento // (raio * 1.4)))
                for i in range(n):
                    t0, t1 = i / n, (i + 1) / n
                    mx = ax + (bx - ax) * (t0 + t1) / 2
                    my = ay + (by - ay) * (t0 + t1) / 2
                    self._preview_ids.append(self.canvas.create_arc(
                        mx - raio, my - raio, mx + raio, my + raio,
                        start=0, extent=180, style="arc", outline=cor, width=largura,
                    ))

            bolhas(x0, y0, x1, y0)
            bolhas(x1, y0, x1, y1)
            bolhas(x1, y1, x0, y1)
            bolhas(x0, y1, x0, y0)
        elif tool == "carimbo":
            item = self.canvas.create_rectangle(x0, y0, x1, y1, outline="#00529B", width=2, dash=(3, 2))
            self._preview_ids.append(item)
        else:
            item = self.canvas.create_rectangle(
                x0, y0, x1, y1, outline=cor, width=largura, dash=dash_canvas or (4, 2)
            )
            self._preview_ids.append(item)

    def var_seta_ativa(self):
        return self._var_seta_linha.get()

    def _limpar_preview(self):
        for item in self._preview_ids:
            self.canvas.delete(item)
        self._preview_ids = []

    # ------------------------- criação de marcações ------------------------- #
    def _criar_marcacao_clique(self, tool, event):
        cx = self.canvas.canvasx(event.x)
        cy = self.canvas.canvasy(event.y)
        px, py = cx / self.zoom, cy / self.zoom  # coordenadas reais no PDF

        marcacao = {"tipo": tool, "pagina": self.pagina_atual, "x": px, "y": py, "data": hoje_str()}
        if tool == "texto_livre":
            marcacao["fontsize"] = self.var_fontsize.get()
            marcacao["cor"] = self.cor_texto_livre

        cor_ativa = self._cor_ativa_da_ferramenta(tool)
        if tool in FERRAMENTAS_SUJEITAS_A_CORRETO and cor_ativa and cor_e_amarela(cor_ativa):
            self._marcar_como_correto(marcacao)
            self.comentarios.append(marcacao)
            self._atualizar_lista_comentarios()
            self.renderizar_pagina()
            return

        kwargs = {}
        if tool == "texto_livre":
            kwargs = {"fontsize": self.var_fontsize.get(), "cor": self.cor_texto_livre}

        dialog = ComentarioDialog(self.root, self.pagina_atual, tipo=tool, **kwargs)
        if not dialog.result:
            return  # marcação descartada se o usuário cancelar

        marcacao["texto"] = dialog.result["texto"]
        marcacao["classificacao"] = dialog.result["classificacao"]
        if tool == "texto_livre":
            marcacao["fontsize"] = dialog.result["fontsize"]
            marcacao["cor"] = dialog.result["cor"]
        self._aplicar_status(marcacao, dialog.result.get("status", "Pendente"))
        self._aplicar_verificacao_correcao(marcacao, dialog.result.get("verificacao_correcao", False))

        self.comentarios.append(marcacao)
        self._atualizar_lista_comentarios()
        self.renderizar_pagina()

    def _criar_marcacao_arrasto(self, tool, cx0, cy0, cx1, cy1):
        x0, y0 = cx0 / self.zoom, cy0 / self.zoom
        x1, y1 = cx1 / self.zoom, cy1 / self.zoom

        marcacao = {
            "tipo": tool, "pagina": self.pagina_atual,
            "x0": x0, "y0": y0, "x1": x1, "y1": y1, "data": hoje_str(),
        }
        if tool in FERRAMENTAS_COM_BORDA:
            marcacao["cor_borda"] = self.estilo_cor_borda
            marcacao["espessura"] = self.estilo_espessura.get()
            marcacao["tracejado"] = self.estilo_tracejado.get()
        if tool in FERRAMENTAS_COM_PREENCHIMENTO:
            marcacao["cor_preenchimento"] = self.estilo_cor_preenchimento
            marcacao["opacidade_borda"] = self.estilo_opacidade_borda.get()
            marcacao["opacidade_preenchimento"] = self.estilo_opacidade_preenchimento.get()
        if tool == "linha":
            marcacao["seta"] = self.var_seta_ativa()
        if tool in ("sublinhado", "tachado"):
            marcacao["cor"] = self.estilo_cor_borda

        cor_ativa = self._cor_ativa_da_ferramenta(tool)
        if tool in FERRAMENTAS_SUJEITAS_A_CORRETO and cor_ativa and cor_e_amarela(cor_ativa):
            self._marcar_como_correto(marcacao)
            self.comentarios.append(marcacao)
            self._atualizar_lista_comentarios()
            self.renderizar_pagina()
            return

        kwargs = {}
        if tool == "realce":
            kwargs = {"cor": CORES_REALCE[self.var_cor_realce.get()]}

        dialog = ComentarioDialog(self.root, self.pagina_atual, tipo=tool, **kwargs)
        if not dialog.result:
            return  # marcação descartada se o usuário cancelar

        marcacao["texto"] = dialog.result["texto"]
        marcacao["classificacao"] = dialog.result["classificacao"]
        if tool == "realce":
            marcacao["cor"] = dialog.result["cor"]
        self._aplicar_status(marcacao, dialog.result.get("status", "Pendente"))
        self._aplicar_verificacao_correcao(marcacao, dialog.result.get("verificacao_correcao", False))

        self.comentarios.append(marcacao)
        self._atualizar_lista_comentarios()
        self.renderizar_pagina()

    def _criar_carimbo(self, event):
        cx = self.canvas.canvasx(event.x)
        cy = self.canvas.canvasy(event.y)
        px, py = cx / self.zoom, cy / self.zoom
        tamanho = (
            CARIMBO_TAMANHO_COM_LOGO
            if (not self.carimbo_imagem_path and os.path.exists(CAMINHO_LOGO_PADRAO))
            else CARIMBO_TAMANHO_PADRAO
        )
        largura, altura = tamanho

        verificador_padrao = self.dados_documento.get("verificado_por", "")
        data_padrao = self.dados_documento.get("data_verificacao", "") or hoje_str()

        dialog = CarimboDialog(self.root, verificador_padrao, data_padrao)
        if not dialog.result:
            return

        verificador = dialog.result["verificador"]
        data_txt = dialog.result["data"]

        marcacao = {
            "tipo": "carimbo",
            "pagina": self.pagina_atual,
            "x0": px, "y0": py, "x1": px + largura, "y1": py + altura,
            "verificador": verificador,
            "data_carimbo": data_txt,
            "imagem_path": self.carimbo_imagem_path,
            "classificacao": "",
            "texto": f"Carimbo de verificação - Verificador: {verificador} - Data: {data_txt}",
            "data": hoje_str(),
        }
        self.comentarios.append(marcacao)
        self._atualizar_lista_comentarios()
        self.renderizar_pagina()

    def _configurar_carimbo(self):
        caminho = filedialog.askopenfilename(
            title="Selecione a imagem do carimbo (PNG recomendado, fundo transparente)",
            filetypes=[("Imagens", "*.png *.jpg *.jpeg"), ("Todos os arquivos", "*.*")],
        )
        if caminho:
            self.carimbo_imagem_path = caminho
            self._atualizar_opcoes_ferramenta()

    def _editar_carimbo(self, idx):
        c = self.comentarios[idx]
        dialog = CarimboDialog(self.root, c.get("verificador", ""), c.get("data_carimbo", ""))
        if dialog.result:
            c["verificador"] = dialog.result["verificador"]
            c["data_carimbo"] = dialog.result["data"]
            c["texto"] = (
                f"Carimbo de verificação - Verificador: {c['verificador']} "
                f"- Data: {c['data_carimbo']}"
            )
            self._atualizar_lista_comentarios()
            self.renderizar_pagina()

    def _atualizar_lista_comentarios(self):
        self.tree.delete(*self.tree.get_children())
        filtro = self.filtro_status_var.get() if hasattr(self, "filtro_status_var") else "Todos"
        contagem = {s: 0 for s in STATUS_REVISAO}

        for idx, c in enumerate(self.comentarios):
            status = c.get("status", "Pendente")
            contagem[status] = contagem.get(status, 0) + 1
            if filtro != "Todos" and status != filtro:
                continue
            resumo = c["texto"][:40] + ("…" if len(c["texto"]) > 40 else "")
            tipo_label = FERRAMENTAS.get(c.get("tipo", "comentario"), {}).get(
                "label", "Comentário"
            )
            tag = STATUS_TAGS.get(status, "status_pendente")
            self.tree.insert(
                "", "end", iid=str(idx), tags=(tag,),
                values=(tipo_label, c["pagina"] + 1, c["classificacao"], resumo, status)
            )

        total = len(self.comentarios)
        self.lbl_status.config(
            text=(
                f"{contagem['Pendente']} pendentes / {contagem['Corrigido']} corrigidos / "
                f"{contagem['Não corrigido']} não corrigidos / "
                f"{contagem['Reincidente']} reincidentes / {contagem['Correto']} corretos / "
                f"{total} total"
            )
        )

    def _indices_selecionados(self):
        return [int(iid) for iid in self.tree.selection()]

    def _indice_selecionado(self):
        indices = self._indices_selecionados()
        return indices[0] if indices else None

    def editar_comentario_selecionado(self, event=None):
        indices = self._indices_selecionados()
        if not indices:
            return
        if len(indices) > 1:
            messagebox.showinfo("Aviso", "Selecione apenas um item para editar.")
            return
        idx = indices[0]
        c = self.comentarios[idx]
        tipo = c.get("tipo", "comentario")

        if tipo == "carimbo":
            self._editar_carimbo(idx)
            return

        kwargs = {
            "status": c.get("status", "Pendente"),
            "verificacao_correcao": c.get("verificacao_correcao", False),
        }
        if tipo == "texto_livre":
            kwargs.update(fontsize=c.get("fontsize", 12), cor=c.get("cor", "#000000"))
        elif tipo == "realce":
            kwargs.update(cor=c.get("cor", "#FFFF00"))

        dialog = ComentarioDialog(
            self.root, c["pagina"], tipo=tipo, texto=c["texto"],
            classificacao=c["classificacao"], **kwargs
        )
        if dialog.result:
            status_anterior = c.get("status", "Pendente")
            verificacao_anterior = c.get("verificacao_correcao", False)
            c["texto"] = dialog.result["texto"]
            c["classificacao"] = dialog.result["classificacao"]
            if tipo == "texto_livre":
                c["fontsize"] = dialog.result["fontsize"]
                c["cor"] = dialog.result["cor"]
            elif tipo == "realce":
                c["cor"] = dialog.result["cor"]
            novo_status = dialog.result.get("status", status_anterior)
            if novo_status != status_anterior:
                self._aplicar_status(c, novo_status)
            novo_verificacao = dialog.result.get("verificacao_correcao", False)
            if novo_verificacao != verificacao_anterior:
                self._aplicar_verificacao_correcao(c, novo_verificacao)
            self._atualizar_lista_comentarios()
            self.renderizar_pagina()

    def excluir_comentario_selecionado(self):
        indices = sorted(self._indices_selecionados(), reverse=True)
        if not indices:
            return
        msg = (
            "Excluir esta marcação?" if len(indices) == 1
            else f"Excluir {len(indices)} itens selecionados?"
        )
        if messagebox.askyesno("Confirmar", msg):
            for idx in indices:
                del self.comentarios[idx]
            self._atualizar_lista_comentarios()
            self.renderizar_pagina()

    def ir_ate_comentario(self):
        idx = self._indice_selecionado()
        if idx is None:
            return
        c = self.comentarios[idx]
        self.pagina_atual = c["pagina"]
        self.renderizar_pagina()

    # ------------------------- salvar/abrir projeto (JSON) ------------------------- #
    def salvar_projeto(self):
        if not self.doc:
            messagebox.showinfo("Aviso", "Abra um PDF primeiro.")
            return

        if not self._avisou_pdf_embutido:
            messagebox.showinfo(
                "Aviso",
                "O arquivo de projeto vai incluir uma cópia do PDF original, para que "
                "possa ser aberto em outro computador sem precisar do arquivo original. "
                "Isso é esperado e pode deixar o arquivo de projeto grande.",
            )
            self._avisou_pdf_embutido = True

        caminho = filedialog.asksaveasfilename(
            defaultextension=".json",
            filetypes=[("Projeto de Revisão (*.json)", "*.json")],
            title="Salvar projeto",
        )
        if not caminho:
            return

        with open(self.pdf_path, "rb") as f:
            pdf_base64 = base64.b64encode(f.read()).decode("ascii")

        dados = {
            "pdf_original": os.path.basename(self.pdf_path),
            "pdf_base64": pdf_base64,
            "dados_documento": self.dados_documento,
            "comentarios": self.comentarios,
        }
        with open(caminho, "w", encoding="utf-8") as f:
            json.dump(dados, f, ensure_ascii=False, indent=2)
        messagebox.showinfo("Concluído", "Projeto salvo com sucesso.")

    def abrir_projeto(self):
        caminho = filedialog.askopenfilename(
            filetypes=[("Projeto de Revisão (*.json)", "*.json")], title="Abrir projeto"
        )
        if not caminho:
            return
        with open(caminho, "r", encoding="utf-8") as f:
            dados = json.load(f)

        pdf_base64 = dados.get("pdf_base64")
        if pdf_base64:
            pdf_bytes = base64.b64decode(pdf_base64)
            arq_tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
            arq_tmp.write(pdf_bytes)
            arq_tmp.close()
            pdf_original = arq_tmp.name
        else:
            # projeto salvo antes da Parte 4 (sem PDF embutido): mantém o comportamento antigo
            pdf_original = dados.get("pdf_original")
            if not pdf_original or not os.path.exists(pdf_original):
                pdf_original = filedialog.askopenfilename(
                    title="Localize o PDF original", filetypes=[("Arquivos PDF", "*.pdf")]
                )
                if not pdf_original:
                    return

        self.doc = fitz.open(pdf_original)
        self.pdf_path = pdf_original
        self.dados_documento = dados.get("dados_documento", {})
        self.comentarios = dados.get("comentarios", [])
        self.pagina_atual = 0
        self._atualizar_lista_comentarios()
        self.renderizar_pagina()
        self.lbl_status.config(text=f"Projeto carregado: {caminho}")

    # ------------------------- exportar PDF ------------------------- #
    def exportar_pdf(self):
        if not self.doc:
            messagebox.showinfo("Aviso", "Abra um PDF primeiro.")
            return
        if not self.comentarios:
            if not messagebox.askyesno(
                "Aviso", "Não há marcações inseridas. Exportar mesmo assim?"
            ):
                return

        destino = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("PDF", "*.pdf")],
            title="Salvar PDF com comentários",
        )
        if not destino:
            return

        doc_saida = fitz.open(self.pdf_path)

        autor = self.dados_documento.get("elaborado_por", "")
        for idx, c in enumerate(self.comentarios, start=1):
            pagina = doc_saida[c["pagina"]]
            tipo = c.get("tipo", "comentario")

            if tipo == "carimbo":
                rect = rect_de_marcacao(c)
                imagem_bytes = gerar_imagem_carimbo(
                    c.get("imagem_path"), c.get("verificador", ""), c.get("data_carimbo", "")
                )
                pagina.insert_image(rect, stream=imagem_bytes)
                continue

            status = c.get("status", "Pendente")
            conteudo = f"[{status.upper()}] [Nº {idx}] {c['classificacao']}\n\n{c['texto']}"
            annot = None

            if tipo == "comentario":
                ponto = fitz.Point(c["x"], c["y"])
                annot = pagina.add_text_annot(ponto, conteudo, icon="Comment")
                annot.set_colors(stroke=hex_para_rgb01(cor_status(c)))
            elif tipo == "texto_livre":
                fontsize = c.get("fontsize", 12)
                rect = rect_texto_livre(c["x"], c["y"], c["texto"], fontsize)
                annot = pagina.add_freetext_annot(
                    rect, c["texto"], fontsize=fontsize,
                    text_color=hex_para_rgb01(c.get("cor", "#000000")),
                )
            elif tipo == "caneta":
                pontos = [(p[0], p[1]) for p in c["pontos"]]
                if len(pontos) < 2:
                    continue
                annot = pagina.add_ink_annot([pontos])
                annot.set_colors(stroke=hex_para_rgb01(c.get("cor", "#FF0000")))
                annot.set_border(width=c.get("espessura", 3))
                annot.set_opacity(max(0.0, min(1.0, c.get("opacidade", 100) / 100)))
            elif tipo == "retangulo":
                annot = pagina.add_rect_annot(rect_de_marcacao(c))
            elif tipo == "elipse":
                annot = pagina.add_circle_annot(rect_de_marcacao(c))
            elif tipo == "nuvem":
                annot = pagina.add_rect_annot(rect_de_marcacao(c))
                annot.set_border(width=espessura_efetiva(c), clouds=2)
            elif tipo == "linha":
                annot = pagina.add_line_annot(
                    fitz.Point(c["x0"], c["y0"]), fitz.Point(c["x1"], c["y1"])
                )
                if c.get("seta"):
                    annot.set_line_ends(fitz.PDF_ANNOT_LE_NONE, fitz.PDF_ANNOT_LE_OPEN_ARROW)
            elif tipo == "seta":
                annot = pagina.add_line_annot(
                    fitz.Point(c["x0"], c["y0"]), fitz.Point(c["x1"], c["y1"])
                )
                annot.set_line_ends(fitz.PDF_ANNOT_LE_NONE, fitz.PDF_ANNOT_LE_OPEN_ARROW)
            elif tipo == "sublinhado":
                annot = pagina.add_underline_annot(rect_de_marcacao(c).quad)
            elif tipo == "tachado":
                annot = pagina.add_strikeout_annot(rect_de_marcacao(c).quad)
            elif tipo == "realce":
                annot = pagina.add_highlight_annot(rect_de_marcacao(c).quad)

            if annot is None:
                continue

            if tipo in FERRAMENTAS_COM_BORDA:
                cor_borda_final = cor_com_opacidade_simulada(
                    c.get("cor_borda", "#FF0000"), c.get("opacidade_borda", 100)
                )
                cores = {"stroke": cor_borda_final}
                if tipo in FERRAMENTAS_COM_PREENCHIMENTO and c.get("cor_preenchimento"):
                    cores["fill"] = cor_com_opacidade_simulada(
                        c["cor_preenchimento"], c.get("opacidade_preenchimento", 100)
                    )
                annot.set_colors(**cores)
                if tipo != "nuvem":
                    annot.set_border(
                        width=espessura_efetiva(c),
                        dashes=DASHES_PDF.get(c.get("tracejado", "Contínua")),
                    )
            elif tipo in ("sublinhado", "tachado"):
                annot.set_colors(stroke=hex_para_rgb01(c.get("cor", "#000000")))
            elif tipo == "realce":
                annot.set_colors(stroke=hex_para_rgb01(c.get("cor", "#FFFF00")))

            annot.set_info(title=autor or "Revisor", subject=c["classificacao"], content=conteudo)
            annot.update()

            if status == "Corrigido":
                rect_conf = rect_confirmacao(c)
                elipse = pagina.add_circle_annot(rect_conf)
                elipse.set_colors(stroke=hex_para_rgb01(COR_AZUL_CONFIRMACAO))
                elipse.set_border(width=1.5)
                elipse.update()

                if c.get("verificacao_correcao"):
                    diagonais = (
                        (fitz.Point(rect_conf.x0, rect_conf.y0), fitz.Point(rect_conf.x1, rect_conf.y1)),
                        (fitz.Point(rect_conf.x1, rect_conf.y0), fitz.Point(rect_conf.x0, rect_conf.y1)),
                    )
                    for p1, p2 in diagonais:
                        linha_x = pagina.add_line_annot(p1, p2)
                        linha_x.set_colors(stroke=hex_para_rgb01(COR_VERDE_VALIDACAO))
                        linha_x.set_border(width=1.5)
                        linha_x.update()

        # grava metadados gerais do documento
        meta = doc_saida.metadata
        meta.update({
            "author": self.dados_documento.get("elaborado_por", meta.get("author", "")),
            "subject": (
                f"Verificado por: {self.dados_documento.get('verificado_por', '')} | "
                f"Data elaboração: {self.dados_documento.get('data_elaboracao', '')} | "
                f"Data verificação: {self.dados_documento.get('data_verificacao', '')}"
            ),
        })
        doc_saida.set_metadata(meta)

        doc_saida.save(destino)
        doc_saida.close()
        messagebox.showinfo("Concluído", f"PDF exportado:\n{destino}")

    # ------------------------- exportar Excel ------------------------- #
    def _linhas_excel(self):
        """Monta uma linha de dados por marcação, no formato exato pedido para a
        aba 'Comentários' (uma linha = uma ocorrência). Carimbo e itens com
        status 'Correto' entram com CODIGO ERRO / QUANTIDADE / TIPO DE ERRO em
        branco, por não representarem um erro. RESPONSÁVEL PELO ERRO é sempre
        o "Elaborado por" do documento (Parte 6, item 13 — não é mais por item)."""
        info = self.dados_documento
        elaborado_por = info.get("elaborado_por", "")
        linhas = []
        for c in self.comentarios:
            tipo = c.get("tipo")
            status = c.get("status", "Pendente")
            is_carimbo = tipo == "carimbo"
            isento_de_erro = is_carimbo or status == "Correto"
            classificacao = c.get("classificacao", "")
            linhas.append({
                "projeto": info.get("projeto", ""),
                "data": c.get("data", ""),
                "codigo": "" if isento_de_erro else CODIGOS_ERRO.get(classificacao, ""),
                "quantidade": "" if isento_de_erro else 1,
                "tipo_erro": "" if isento_de_erro else classificacao,
                "responsavel_erro": elaborado_por,
                "responsavel_verificacao": (
                    c.get("verificador", "") if is_carimbo else info.get("verificado_por", "")
                ),
                "numero_doc": info.get("codigo_doc", ""),
                "status": status,
                "revisado_por": c.get("revisado_por", ""),
                "data_revisao": c.get("data_revisao", ""),
                "verificacao_correcao": "Sim" if c.get("verificacao_correcao") else "Não",
                "verificado_correcao_por": c.get("verificado_correcao_por", ""),
                "correcao_verificada": bool(status == "Corrigido" and c.get("verificacao_correcao")),
                "classificacao_raw": classificacao,
                "is_carimbo": is_carimbo,
            })
        return linhas

    def exportar_excel(self):
        if not self.comentarios:
            messagebox.showinfo("Aviso", "Não há marcações para exportar.")
            return

        destino = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            title="Exportar Excel",
        )
        if not destino:
            return

        wb = openpyxl.Workbook()

        cabecalho_font = Font(bold=True, color="FFFFFF")
        cabecalho_fill = PatternFill("solid", fgColor="305496")
        borda = Border(*(Side(style="thin"),) * 4)

        # ---------- Aba 1: Comentários (uma linha por ocorrência) ---------- #
        ws = wb.active
        ws.title = "Comentários"

        colunas = [
            "PROJETO", "DATA", "CODIGO ERRO", "QUANTIDADE", "TIPO DE ERRO",
            "RESPONSÁVEL PELO ERRO", "RESPONSÁVEL PELA VERIFICAÇÃO", "NÚMERO DO DOCUMENTO",
            "STATUS", "REVISADO POR", "DATA DA REVISÃO",
            "VERIFICAÇÃO DA CORREÇÃO", "VERIFICADO POR (CORREÇÃO)",
        ]
        for j, titulo in enumerate(colunas, start=1):
            cel = ws.cell(row=1, column=j, value=titulo)
            cel.font = cabecalho_font
            cel.fill = cabecalho_fill
            cel.alignment = Alignment(horizontal="center", vertical="center")
            cel.border = borda

        linhas_dados = self._linhas_excel()
        for i, linha in enumerate(linhas_dados, start=2):
            valores = [
                linha["projeto"], linha["data"], linha["codigo"], linha["quantidade"],
                linha["tipo_erro"], linha["responsavel_erro"],
                linha["responsavel_verificacao"], linha["numero_doc"],
                linha["status"], linha["revisado_por"], linha["data_revisao"],
                linha["verificacao_correcao"], linha["verificado_correcao_por"],
            ]
            for j, v in enumerate(valores, start=1):
                cel = ws.cell(row=i, column=j, value=v)
                cel.border = borda
                cel.alignment = Alignment(vertical="top", wrap_text=(j == 5))

        larguras = [22, 12, 12, 11, 42, 24, 26, 20, 14, 18, 14, 20, 22]
        for j, largura in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(j)].width = largura

        # ---------- Aba 2: Resumo (agrupa por CODIGO ERRO, soma QUANTIDADE) ---------- #
        ws2 = wb.create_sheet("Resumo")
        ws2.append(["CODIGO ERRO", "TIPO DE ERRO", "QUANTIDADE"])
        for j in range(1, 4):
            cel = ws2.cell(row=1, column=j)
            cel.font = cabecalho_font
            cel.fill = cabecalho_fill
            cel.alignment = Alignment(horizontal="center")
            cel.border = borda

        contagem = {}
        for linha in linhas_dados:
            if linha["is_carimbo"] or not linha["codigo"]:
                continue
            contagem[linha["classificacao_raw"]] = (
                contagem.get(linha["classificacao_raw"], 0) + (linha["quantidade"] or 0)
            )

        categorias_ordenadas = sorted(
            CLASSIFICACOES, key=lambda k: contagem.get(k, 0), reverse=True
        )

        linha_n = 2
        for cat in categorias_ordenadas:
            qtd = contagem.get(cat, 0)
            ws2.cell(row=linha_n, column=1, value=CODIGOS_ERRO[cat]).border = borda
            ws2.cell(row=linha_n, column=2, value=cat).border = borda
            ws2.cell(row=linha_n, column=3, value=qtd).border = borda
            linha_n += 1

        total_row = linha_n
        ws2.cell(row=total_row, column=2, value="TOTAL").font = Font(bold=True)
        ws2.cell(row=total_row, column=3, value=sum(contagem.values())).font = Font(bold=True)

        ws2.column_dimensions["A"].width = 12
        ws2.column_dimensions["B"].width = 55
        ws2.column_dimensions["C"].width = 14

        chart = BarChart()
        chart.title = "Quantidade de erros por classificação"
        chart.y_axis.title = "Quantidade"
        chart.x_axis.title = "Classificação"
        data_ref = Reference(ws2, min_col=3, min_row=1, max_row=total_row - 1)
        cats_ref = Reference(ws2, min_col=2, min_row=2, max_row=total_row - 1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.width = 26
        chart.height = 12
        ws2.add_chart(chart, "E2")

        # ---------- Aba 3: Painel de Status ---------- #
        ws3 = wb.create_sheet("Painel de Status")
        ws3.append(["Status", "Quantidade"])
        for j in (1, 2):
            cel = ws3.cell(row=1, column=j)
            cel.font = cabecalho_font
            cel.fill = cabecalho_fill
            cel.alignment = Alignment(horizontal="center")
            cel.border = borda

        contagem_status = {s: 0 for s in STATUS_REVISAO}
        for linha in linhas_dados:
            contagem_status[linha["status"]] = contagem_status.get(linha["status"], 0) + 1

        r = 2
        for status in STATUS_REVISAO:
            ws3.cell(row=r, column=1, value=status).border = borda
            ws3.cell(row=r, column=2, value=contagem_status[status]).border = borda
            r += 1
        ws3.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
        ws3.cell(row=r, column=2, value=len(linhas_dados)).font = Font(bold=True)
        r += 1

        correcoes_verificadas = sum(1 for linha in linhas_dados if linha["correcao_verificada"])
        ws3.cell(row=r, column=1, value="Correções verificadas").border = borda
        ws3.cell(row=r, column=2, value=correcoes_verificadas).border = borda

        ws3.column_dimensions["A"].width = 20
        ws3.column_dimensions["B"].width = 14

        wb.save(destino)
        messagebox.showinfo("Concluído", f"Excel exportado:\n{destino}")


def main():
    root = tk.Tk()

    if sv_ttk is not None:
        try:
            sv_ttk.set_theme("light")
        except Exception:
            print("Aviso: falha ao aplicar o tema sv_ttk; seguindo com o tema padrão do sistema.")
    else:
        print("Dica: rode 'pip install sv-ttk' para um visual estilo Windows 11.")

    import tkinter.font as tkfont
    familias = set(tkfont.families())
    nome_fonte = "Segoe UI Variable" if "Segoe UI Variable" in familias else "Segoe UI"
    fonte_padrao = (nome_fonte, 10)
    root.option_add("*Font", fonte_padrao)
    style = ttk.Style()
    style.configure(".", font=fonte_padrao)
    if sv_ttk is None:
        style.configure(".", background="#F3F3F3")
    root.configure(bg="#F3F3F3")

    app = RevisorPDFApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
