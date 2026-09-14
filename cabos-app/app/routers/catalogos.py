import io
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import Response
from sqlalchemy.orm import Session
import openpyxl

from .. import models, schemas
from ..database import get_db

router = APIRouter(prefix="/api/catalogos", tags=["catalogos"])
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _serialize_cabo(c: models.CatalogoCabo):
    return {
        "id": c.id, "fabricante": c.fabricante, "linha_produto": c.linha_produto, "aplicacao": c.aplicacao,
        "material_condutor": c.material_condutor, "tipo_isolacao": c.tipo_isolacao,
        "possui_blindagem": c.possui_blindagem, "tensao_isolamento": c.tensao_isolamento,
        "construcao_basica": c.construcao_basica, "num_condutores": c.num_condutores,
        "secao_nominal_mm2": c.secao_nominal_mm2, "diametro_externo_nominal_mm": c.diametro_externo_nominal_mm,
        "peso_kg_km": c.peso_kg_km, "resistencia_condutor_20c_ohm_km": c.resistencia_condutor_20c_ohm_km,
        "reatancia_ohm_km": c.reatancia_ohm_km, "capacidade_conducao_a": c.capacidade_conducao_a,
        "norma_referencia": c.norma_referencia, "fonte_datasheet": c.fonte_datasheet, "area_secao_mm2": c.area_secao_mm2,
    }


def _serialize_infra(i: models.CatalogoInfraestrutura):
    return {
        "id": i.id, "fabricante": i.fabricante, "linha_produto": i.linha_produto, "tipo": i.tipo,
        "diametro_nominal_pol": i.diametro_nominal_pol, "diametro_nominal_mm": i.diametro_nominal_mm,
        "parede_mm": i.parede_mm, "diametro_externo_mm": i.diametro_externo_mm,
        "diametro_interno_mm": i.diametro_interno_mm, "area_util_mm2": i.area_util_mm2,
        "largura_nominal_mm": i.largura_nominal_mm, "altura_nominal_mm": i.altura_nominal_mm,
        "largura_util_mm": i.largura_util_mm, "fonte_datasheet": i.fonte_datasheet, "customizado": i.customizado,
    }


# ---------------- Catálogo de cabos ----------------

@router.get("/cabos")
def listar_catalogo_cabos(
    tipo_isolacao: Optional[str] = None, linha_produto: Optional[str] = None,
    secao_min: Optional[float] = None, q: Optional[str] = None, db: Session = Depends(get_db),
):
    query = db.query(models.CatalogoCabo)
    if tipo_isolacao:
        query = query.filter(models.CatalogoCabo.tipo_isolacao == tipo_isolacao)
    if linha_produto:
        query = query.filter(models.CatalogoCabo.linha_produto == linha_produto)
    if secao_min is not None:
        query = query.filter(models.CatalogoCabo.secao_nominal_mm2 >= secao_min)
    if q:
        like = f"%{q}%"
        query = query.filter(models.CatalogoCabo.linha_produto.ilike(like) | models.CatalogoCabo.fabricante.ilike(like))
    itens = query.order_by(models.CatalogoCabo.linha_produto, models.CatalogoCabo.num_condutores, models.CatalogoCabo.secao_nominal_mm2).all()
    return [_serialize_cabo(c) for c in itens]


@router.post("/cabos")
def criar_catalogo_cabo(payload: schemas.CatalogoCaboIn, db: Session = Depends(get_db)):
    item = models.CatalogoCabo(**payload.model_dump())
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize_cabo(item)


@router.put("/cabos/{item_id}")
def atualizar_catalogo_cabo(item_id: int, payload: schemas.CatalogoCaboIn, db: Session = Depends(get_db)):
    item = db.get(models.CatalogoCabo, item_id)
    if not item:
        raise HTTPException(404, "Item não encontrado.")
    for k, v in payload.model_dump().items():
        setattr(item, k, v)
    db.commit()
    db.refresh(item)
    return _serialize_cabo(item)


@router.delete("/cabos/{item_id}")
def remover_catalogo_cabo(item_id: int, db: Session = Depends(get_db)):
    item = db.get(models.CatalogoCabo, item_id)
    if not item:
        raise HTTPException(404, "Item não encontrado.")
    em_uso = db.query(models.Cabo).filter(models.Cabo.catalogo_cabo_id == item_id).count()
    if em_uso:
        raise HTTPException(409, f"Este cabo do catálogo é usado por {em_uso} cabo(s) de projeto.")
    db.delete(item)
    db.commit()
    return {"ok": True}


# ---------------- Catálogo de infraestrutura ----------------

@router.get("/infraestrutura")
def listar_catalogo_infra(
    tipo: Optional[str] = None, linha_produto: Optional[str] = None, q: Optional[str] = None, db: Session = Depends(get_db),
):
    query = db.query(models.CatalogoInfraestrutura)
    if tipo:
        query = query.filter(models.CatalogoInfraestrutura.tipo == tipo)
    if linha_produto:
        query = query.filter(models.CatalogoInfraestrutura.linha_produto == linha_produto)
    if q:
        like = f"%{q}%"
        query = query.filter(models.CatalogoInfraestrutura.linha_produto.ilike(like))
    itens = query.order_by(models.CatalogoInfraestrutura.tipo, models.CatalogoInfraestrutura.linha_produto,
                           models.CatalogoInfraestrutura.diametro_nominal_mm, models.CatalogoInfraestrutura.largura_nominal_mm).all()
    return [_serialize_infra(i) for i in itens]


@router.post("/infraestrutura/customizada")
def criar_infra_customizada(payload: dict, db: Session = Depends(get_db)):
    item = models.CatalogoInfraestrutura(
        fabricante=payload.get("fabricante", "Customizado"),
        linha_produto=payload["linha_produto"], tipo=payload["tipo"],
        largura_nominal_mm=payload.get("largura_nominal_mm"), altura_nominal_mm=payload.get("altura_nominal_mm"),
        largura_util_mm=payload.get("largura_util_mm") or payload.get("largura_nominal_mm"),
        fonte_datasheet="Cadastro avulso do usuário", customizado=True,
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize_infra(item)


# ---------------- Tabelas normativas (somente leitura) ----------------

@router.get("/normas/capacidade-conducao")
def tab_capacidade(db: Session = Depends(get_db)):
    rows = db.query(models.TabCapacidadeConducao).order_by(
        models.TabCapacidadeConducao.isolacao_grupo, models.TabCapacidadeConducao.metodo_instalacao,
        models.TabCapacidadeConducao.num_condutores_carregados, models.TabCapacidadeConducao.secao_mm2).all()
    return [{"isolacao_grupo": r.isolacao_grupo, "metodo_instalacao": r.metodo_instalacao,
             "num_condutores_carregados": r.num_condutores_carregados, "secao_mm2": r.secao_mm2,
             "capacidade_a": r.capacidade_a} for r in rows]


@router.get("/normas/fator-temperatura")
def tab_temperatura(db: Session = Depends(get_db)):
    rows = db.query(models.TabFatorTemperatura).order_by(models.TabFatorTemperatura.isolacao_grupo, models.TabFatorTemperatura.temperatura_c).all()
    return [{"isolacao_grupo": r.isolacao_grupo, "temperatura_c": r.temperatura_c, "fator": r.fator} for r in rows]


@router.get("/normas/fator-agrupamento")
def tab_agrupamento(db: Session = Depends(get_db)):
    rows = db.query(models.TabFatorAgrupamento).order_by(models.TabFatorAgrupamento.num_circuitos).all()
    return [{"num_circuitos": r.num_circuitos, "fator": r.fator} for r in rows]


@router.get("/normas/resistividade")
def tab_resistividade(db: Session = Depends(get_db)):
    rows = db.query(models.TabResistividadeCondutor).order_by(models.TabResistividadeCondutor.secao_mm2).all()
    return [{"secao_mm2": r.secao_mm2, "resistencia_ohm_km": r.resistencia_ohm_km} for r in rows]


@router.get("/normas/limite-ocupacao-eletroduto")
def tab_limite_eletroduto(db: Session = Depends(get_db)):
    rows = db.query(models.TabLimiteOcupacaoEletroduto).all()
    return [{"num_cabos": r.num_cabos, "taxa_max_pct": r.taxa_max_pct} for r in rows]


# ---------------- Importação / modelos de planilha ----------------

CAMPOS_CABO = ["fabricante", "linha_produto", "aplicacao", "material_condutor", "tipo_isolacao",
               "possui_blindagem", "tensao_isolamento", "construcao_basica", "num_condutores",
               "secao_nominal_mm2", "diametro_externo_nominal_mm", "peso_kg_km",
               "resistencia_condutor_20c_ohm_km", "reatancia_ohm_km", "capacidade_conducao_a",
               "norma_referencia", "fonte_datasheet"]

CAMPOS_INFRA = ["fabricante", "linha_produto", "tipo", "diametro_nominal_pol", "diametro_nominal_mm",
                "parede_mm", "diametro_externo_mm", "diametro_interno_mm", "area_util_mm2",
                "largura_nominal_mm", "altura_nominal_mm", "largura_util_mm", "fonte_datasheet"]

CAMPOS_CAPACIDADE = ["isolacao_grupo", "metodo_instalacao", "num_condutores_carregados", "secao_mm2", "capacidade_a"]


def _planilha_modelo(campos, exemplo=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "dados"
    ws.append(campos)
    if exemplo:
        ws.append([exemplo.get(c) for c in campos])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/cabos/modelo.xlsx")
def modelo_cabos():
    exemplo = {"fabricante": "Prysmian", "linha_produto": "Afumex Green 1kV", "aplicacao": "energia BT",
               "material_condutor": "cobre", "tipo_isolacao": "XLPE", "possui_blindagem": False,
               "tensao_isolamento": "0,6/1 kV", "construcao_basica": "3x2.5", "num_condutores": 3,
               "secao_nominal_mm2": 2.5, "diametro_externo_nominal_mm": 10.1, "peso_kg_km": 120,
               "resistencia_condutor_20c_ohm_km": 7.41, "reatancia_ohm_km": 0.09, "capacidade_conducao_a": None,
               "norma_referencia": "NBR 13248", "fonte_datasheet": "datasheet.pdf"}
    return Response(_planilha_modelo(CAMPOS_CABO, exemplo), media_type=XLSX_MEDIA,
                    headers={"Content-Disposition": 'attachment; filename="modelo_catalogo_cabos.xlsx"'})


@router.get("/infraestrutura/modelo.xlsx")
def modelo_infra():
    exemplo = {"fabricante": "Elecon", "linha_produto": "Eletroduto Rígido Médio", "tipo": "eletroduto",
               "diametro_nominal_pol": '1"', "diametro_nominal_mm": 25, "parede_mm": 0.9, "diametro_externo_mm": 31.9,
               "diametro_interno_mm": 30.1, "area_util_mm2": 711.6, "largura_nominal_mm": None,
               "altura_nominal_mm": None, "largura_util_mm": None, "fonte_datasheet": "Catálogo Elecon 2018"}
    return Response(_planilha_modelo(CAMPOS_INFRA, exemplo), media_type=XLSX_MEDIA,
                    headers={"Content-Disposition": 'attachment; filename="modelo_catalogo_infraestrutura.xlsx"'})


@router.get("/normas/capacidade-conducao/modelo.xlsx")
def modelo_capacidade(db: Session = Depends(get_db)):
    # Exportar TODOS os dados atuais da tabela de capacidade de condução
    rows = db.query(models.TabCapacidadeConducao).order_by(
        models.TabCapacidadeConducao.isolacao_grupo, models.TabCapacidadeConducao.metodo_instalacao,
        models.TabCapacidadeConducao.num_condutores_carregados, models.TabCapacidadeConducao.secao_mm2).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Capacidade"
    ws.append(CAMPOS_CAPACIDADE)
    for r in rows:
        ws.append([r.isolacao_grupo, r.metodo_instalacao, r.num_condutores_carregados, r.secao_mm2, r.capacidade_a])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(buf.getvalue(), media_type=XLSX_MEDIA,
                    headers={"Content-Disposition": 'attachment; filename="modelo_capacidade_conducao.xlsx"'})


def _importar_planilha(arquivo_bytes, campos, modelo, db, substituir, pos_processar=None):
    wb = openpyxl.load_workbook(arquivo_bytes, data_only=True)
    ws = wb.active
    header = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    presentes = [c for c in campos if c in header]
    if not presentes:
        raise HTTPException(400, "Cabeçalho da planilha não corresponde ao esperado. Baixe a planilha-modelo na tela de Configurações.")

    if substituir:
        db.query(modelo).delete()

    count = 0
    for row in ws.iter_rows(min_row=2):
        valores = {header[i]: cell.value for i, cell in enumerate(row) if i < len(header)}
        if not any(v not in (None, "") for v in valores.values()):
            continue
        dados = {campo: valores.get(campo) for campo in presentes}
        for k, v in list(dados.items()):
            if isinstance(v, str) and v.strip().lower() in ("true", "sim", "verdadeiro", "x"):
                dados[k] = True
            elif isinstance(v, str) and v.strip().lower() in ("false", "não", "nao", "falso"):
                dados[k] = False
        if pos_processar:
            dados = pos_processar(dados)
        db.add(modelo(**dados))
        count += 1
    db.commit()
    return count


def _completar_infra(dados):
    ext, parede, interno = dados.get("diametro_externo_mm"), dados.get("parede_mm"), dados.get("diametro_interno_mm")
    if interno in (None, "") and ext and parede:
        interno = float(ext) - 2 * float(parede)
        dados["diametro_interno_mm"] = round(interno, 2)
    if dados.get("area_util_mm2") in (None, "") and interno:
        dados["area_util_mm2"] = round(3.14159265 * (float(interno) / 2) ** 2, 1)
    if dados.get("largura_util_mm") in (None, "") and dados.get("largura_nominal_mm"):
        dados["largura_util_mm"] = dados["largura_nominal_mm"]
    return dados


@router.post("/cabos/importar")
async def importar_catalogo_cabos(substituir: bool = False, arquivo: UploadFile = File(...), db: Session = Depends(get_db)):
    count = _importar_planilha(arquivo.file, CAMPOS_CABO, models.CatalogoCabo, db, substituir)
    return {"importados": count}


@router.post("/infraestrutura/importar")
async def importar_catalogo_infra(substituir: bool = False, arquivo: UploadFile = File(...), db: Session = Depends(get_db)):
    count = _importar_planilha(arquivo.file, CAMPOS_INFRA, models.CatalogoInfraestrutura, db, substituir, _completar_infra)
    return {"importados": count}


@router.post("/normas/capacidade-conducao/importar")
async def importar_capacidade(substituir: bool = True, arquivo: UploadFile = File(...), db: Session = Depends(get_db)):
    count = _importar_planilha(arquivo.file, CAMPOS_CAPACIDADE, models.TabCapacidadeConducao, db, substituir)
    return {"importados": count}
