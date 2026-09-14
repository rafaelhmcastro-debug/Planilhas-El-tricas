import io
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from . import models, calculations

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBTOTAL_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
PAINEL_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
BOLD = Font(bold=True)


def _write_header(ws, headers, row=1):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _autofit(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _to_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _r(v, casas=2):
    return round(v, casas) if isinstance(v, (int, float)) else v


def _fases_txt(num_fases, neutro, terra):
    s = f"{num_fases}F" if num_fases else "-"
    if neutro:
        s += "+N"
    if terra:
        s += "+PE"
    return s


# ---------------------------------------------------------------------------

def exportar_equipamentos(db: Session, projeto: models.Projeto) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Equipamentos"
    headers = ["TAG", "Descrição", "Tipo de carga", "Potência informada", "Unidade", "Rendimento η",
               "P (kW)", "Q (kvar)", "S (kVA)", "Tensão (V)", "Fases", "cos φ", "FD", "In (A)", "Pd (kW)",
               "Painel/Transformador", "Folha de dados (nº)", "Diagrama elétrico (nº)", "Fornecedor", "Observação"]
    _write_header(ws, headers)
    equipamentos = db.query(models.Equipamento).filter(models.Equipamento.projeto_id == projeto.id).order_by(models.Equipamento.tag).all()
    for r, e in enumerate(equipamentos, start=2):
        vals = [e.tag, e.descricao, e.tipo_carga, e.potencia_valor, e.potencia_unidade, e.rendimento,
                _r(e.potencia_ativa_calc_kw), _r(e.potencia_reativa_calc_kvar), _r(e.potencia_aparente_calc_kva),
                e.tensao_v, _fases_txt(e.num_fases, e.possui_neutro, e.possui_terra), e.fator_potencia, e.fator_demanda,
                _r(e.corrente_nominal_a), _r(e.potencia_demanda_kw), e.painel_transformador_tag, e.folha_dados_numero,
                e.diagrama_eletrico_numero, e.fornecedor, e.observacao]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
    _autofit(ws, [14, 26, 14, 12, 9, 11, 10, 10, 10, 10, 10, 8, 8, 10, 10, 18, 16, 18, 16, 26])
    return _to_bytes(wb)


def exportar_cabos(db: Session, projeto: models.Projeto) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Lista de Cabos"
    headers = ["Item", "Identificação (TAG)", "Tensão (V)", "Configuração", "Construção", "Nº Cabos", "Nº Condutores",
               "Seção (mm²)", "Isolamento", "Cabo (catálogo)", "De", "Para", "Percurso", "Distância (m)",
               "Ib (A)", "Queda de Tensão (%)", "Método", "Seção sugerida (mm²)", "Observação"]
    _write_header(ws, headers)
    cabos = db.query(models.Cabo).filter(models.Cabo.projeto_id == projeto.id).order_by(models.Cabo.tag).all()
    for r, c in enumerate(cabos, start=2):
        percurso = " -> ".join(t.eletroduto_bandeja.tag for t in sorted(c.trechos, key=lambda t: t.ordem))
        cat = c.catalogo_cabo
        multipolar = (c.tipo_construcao or "multipolar") == "multipolar"
        if multipolar:
            num_cabos = c.num_cabos_paralelo
            num_cond = cat.num_condutores if cat else c.num_condutores_calc
            construcao = f"Multipolar ({num_cond} vias)"
        else:
            num_cabos = (c.num_condutores_calc or 0) * (c.num_cabos_paralelo or 1)
            num_cond = 1
            construcao = "Singelo (unipolar)"
        vals = [r - 1, c.tag, c.tensao_v, _fases_txt(c.num_fases, None, None) if not cat else None, construcao,
                num_cabos, num_cond, c.secao_mm2, c.tipo_isolacao,
                f"{cat.fabricante} {cat.linha_produto} {cat.construcao_basica}" if cat else "",
                c.de_tag, c.para_tag, percurso, _r(c.distancia_total_m, 1), _r(c.corrente_projeto_a),
                _r(c.queda_tensao_calc_pct), c.metodo_instalacao, c.secao_sugerida_mm2,
                (("ATENÇÃO: seção menor que a sugerida. " if c.alerta_secao_insuficiente else "") + (c.observacao or "")).strip()]
        # configuração de fases (com N/PE) a partir do número de condutores calculado
        carga = calculations._dados_da_carga(c)
        if carga:
            vals[3] = _fases_txt(carga["num_fases"], carga["possui_neutro"], carga["possui_terra"])
        for col, v in enumerate(vals, start=1):
            ws.cell(row=r, column=col, value=v)
    _autofit(ws, [6, 16, 10, 12, 18, 9, 12, 10, 10, 30, 14, 14, 30, 12, 10, 14, 8, 12, 30])
    return _to_bytes(wb)


def exportar_infraestrutura(db: Session, projeto: models.Projeto) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Eletrodutos e Bandejas"
    headers = ["TAG", "Tipo", "Item do catálogo adotado", "DN / Largura nominal (mm)", "Área útil (mm²) / Largura útil (mm)",
               "De", "Para", "Comprimento (m)", "Nº cabos físicos", "Área ocupada (mm²) / Soma Ø (mm)",
               "Ocupação (%)", "Limite (%)", "Item sugerido pelo cálculo", "Situação", "Cabos no trecho"]
    _write_header(ws, headers)
    itens = db.query(models.EletrodutoBandeja).filter(models.EletrodutoBandeja.projeto_id == projeto.id).order_by(models.EletrodutoBandeja.tag).all()
    for r, it in enumerate(itens, start=2):
        res = calculations.calcular_ocupacao(db, it)
        cat = it.catalogo_item
        if it.tipo == "eletroduto":
            dn = cat.diametro_nominal_mm if cat else it.diametro_nominal_mm
            util = it.area_util_mm2
            ocupado = _r(res.get("area_ocupada_mm2"), 1)
        else:
            dn = cat.largura_nominal_mm if cat else None
            util = it.largura_util_mm
            ocupado = _r(res.get("soma_diametros_mm"), 1)
        situacao = "SEM CABOS" if not res.get("num_cabos_fisicos") else ("ACIMA DO LIMITE" if res.get("alerta") else "OK")
        if it.dimensao_definida_manualmente:
            situacao += " (dimensão manual)"
        vals = [it.tag, it.tipo, res.get("item_escolhido") or "(não definido)", dn, util, it.no_origem, it.no_destino,
                it.comprimento_m, res.get("num_cabos_fisicos"), ocupado, _r(res.get("ocupacao_pct"), 1),
                res.get("limite_pct"), res.get("dimensao_sugerida_desc") or "-", situacao,
                ", ".join(f"{c['tag']} ({c['num_cabos_fisicos']}x Ø{_r(c['diametro_externo_mm'], 1)}mm)" for c in res.get("cabos", []))]
        for col, v in enumerate(vals, start=1):
            ws.cell(row=r, column=col, value=v)
    _autofit(ws, [14, 12, 36, 14, 16, 14, 14, 12, 10, 16, 10, 10, 36, 20, 50])
    return _to_bytes(wb)


def exportar_cargas_demanda(db: Session, projeto: models.Projeto) -> bytes:
    """Planilha hierárquica: cada painel lista suas cargas e os painéis a jusante, com subtotais
    parciais e total, até o painel TOP."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Cargas e Demanda"
    headers = ["Nível", "Painel / Carga", "Descrição", "Tipo", "Tensão (V)", "Fases", "P inst. (kW)", "Q inst. (kvar)",
               "FD / F. diversidade", "Pd (kW)", "Qd (kvar)", "Sd (kVA)", "cos φ", "I (A)", "Alimentado por"]
    _write_header(ws, headers)

    paineis = db.query(models.PainelTransformador).filter(models.PainelTransformador.projeto_id == projeto.id).order_by(models.PainelTransformador.tag).all()
    equipamentos = db.query(models.Equipamento).filter(models.Equipamento.projeto_id == projeto.id).order_by(models.Equipamento.tag).all()
    por_painel = {}
    for e in equipamentos:
        if e.painel_transformador_tag:
            por_painel.setdefault(e.painel_transformador_tag.strip(), []).append(e)
    tags = {p.tag for p in paineis}
    filhos = {}
    for p in paineis:
        pai = (p.painel_alimentador_tag or "").strip()
        if pai in tags:
            filhos.setdefault(pai, []).append(p)
    raizes = [p for p in paineis if (p.painel_alimentador_tag or "").strip() not in tags]

    row = [2]

    def escreve(vals, fill=None, bold=False):
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=row[0], column=col, value=v)
            if fill:
                c.fill = fill
            if bold:
                c.font = BOLD
        row[0] += 1

    def painel_bloco(p, nivel, visitados):
        ind = "    " * nivel
        escreve([nivel, f"{ind}{p.tag}", f"Painel {p.tipo}", p.tipo, p.tensao_v, f"{p.num_fases or 3}F", None, None, None,
                 None, None, None, None, None, p.painel_alimentador_tag or "TOP"], fill=PAINEL_FILL, bold=True)
        sp = sq = spd = sqd = si = 0.0
        for e in por_painel.get(p.tag, []):
            escreve([nivel + 1, f"{ind}    {e.tag}", e.descricao, e.tipo_carga, e.tensao_v,
                     _fases_txt(e.num_fases, e.possui_neutro, e.possui_terra),
                     _r(e.potencia_ativa_calc_kw), _r(e.potencia_reativa_calc_kvar), e.fator_demanda,
                     _r(e.potencia_demanda_kw), _r(e.potencia_demanda_reativa_kvar), _r(e.potencia_aparente_calc_kva),
                     e.fator_potencia, _r(e.corrente_nominal_a), p.tag])
            sp += e.potencia_ativa_calc_kw or 0
            sq += e.potencia_reativa_calc_kvar or 0
            spd += e.potencia_demanda_kw or 0
            sqd += e.potencia_demanda_reativa_kvar or 0
            si += e.corrente_nominal_a or 0
        if por_painel.get(p.tag):
            escreve([nivel + 1, f"{ind}    Subtotal cargas diretas de {p.tag}", None, None, None, None, _r(sp), _r(sq), None,
                     _r(spd), _r(sqd), None, None, _r(si), None], fill=SUBTOTAL_FILL, bold=True)
        for f in filhos.get(p.tag, []):
            if f.tag in visitados:
                continue
            painel_bloco(f, nivel + 1, visitados | {f.tag})
            escreve([nivel + 1, f"{ind}    Contribuição do painel {f.tag} em {p.tag}", None, None, None, None,
                     _r(f.potencia_instalada_kw), _r(f.potencia_reativa_instalada_kvar), f.fator_diversidade,
                     _r(f.demanda_total_kw), _r(f.demanda_reativa_kvar), _r(f.demanda_aparente_kva),
                     _r(f.fator_potencia_calc), _r(f.corrente_total_a), p.tag], fill=SUBTOTAL_FILL)
        escreve([nivel, f"{ind}TOTAL {p.tag}" + (" (TOP)" if nivel == 0 else ""), "demanda com fator de diversidade", p.tipo,
                 p.tensao_v, f"{p.num_fases or 3}F", _r(p.potencia_instalada_kw), _r(p.potencia_reativa_instalada_kvar),
                 p.fator_diversidade, _r(p.demanda_total_kw), _r(p.demanda_reativa_kvar), _r(p.demanda_aparente_kva),
                 _r(p.fator_potencia_calc), _r(p.corrente_total_a), p.painel_alimentador_tag or "TOP"], fill=TOTAL_FILL, bold=True)
        row[0] += 1

    for p in raizes:
        painel_bloco(p, 0, {p.tag})

    orfaos = [e for e in equipamentos if not (e.painel_transformador_tag or "").strip()]
    if orfaos:
        escreve(["", "Equipamentos sem painel informado"], fill=PAINEL_FILL, bold=True)
        for e in orfaos:
            escreve(["", e.tag, e.descricao, e.tipo_carga, e.tensao_v, _fases_txt(e.num_fases, e.possui_neutro, e.possui_terra),
                     _r(e.potencia_ativa_calc_kw), _r(e.potencia_reativa_calc_kvar), e.fator_demanda,
                     _r(e.potencia_demanda_kw), _r(e.potencia_demanda_reativa_kvar), _r(e.potencia_aparente_calc_kva),
                     e.fator_potencia, _r(e.corrente_nominal_a), "-"])
    _autofit(ws, [6, 40, 30, 12, 10, 10, 12, 12, 12, 10, 10, 10, 8, 10, 16])

    ws2 = wb.create_sheet("Resumo por painel")
    _write_header(ws2, ["Painel", "Tipo", "Alimentado por", "Tensão (V)", "P inst. (kW)", "Q inst. (kvar)",
                        "F. diversidade", "Demanda (kW)", "Demanda (kvar)", "Demanda (kVA)", "cos φ", "I alimentador (A)", "Σ In cargas (A)"])
    for r, p in enumerate(paineis, start=2):
        vals = [p.tag, p.tipo, p.painel_alimentador_tag or "TOP", p.tensao_v, _r(p.potencia_instalada_kw),
                _r(p.potencia_reativa_instalada_kvar), p.fator_diversidade, _r(p.demanda_total_kw), _r(p.demanda_reativa_kvar),
                _r(p.demanda_aparente_kva), _r(p.fator_potencia_calc), _r(p.corrente_total_a), _r(p.corrente_instalada_a)]
        for col, v in enumerate(vals, start=1):
            ws2.cell(row=r, column=col, value=v)
    _autofit(ws2, [16, 12, 16, 10, 12, 12, 12, 12, 12, 12, 8, 14, 14])
    return _to_bytes(wb)


def exportar_memoria_calculo(db: Session, projeto: models.Projeto, cabo_id: int = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    query = db.query(models.Cabo).filter(models.Cabo.projeto_id == projeto.id)
    if cabo_id:
        query = query.filter(models.Cabo.id == cabo_id)
    cabos = query.order_by(models.Cabo.tag).all()
    for c in cabos:
        titulo = "".join(ch for ch in c.tag if ch not in '[]:*?/\\')[:28] or f"Cabo{c.id}"
        ws = wb.create_sheet(titulo)
        ws.cell(row=1, column=1, value=f"Memória de cálculo — cabo {c.tag} (de {c.de_tag} para {c.para_tag})").font = Font(bold=True, size=13)
        ws.cell(row=2, column=1, value=f"Projeto {projeto.numero_projeto} — {projeto.nome_projeto} — {projeto.revisao}")
        row = 4
        if c.memoria_calculo_json:
            memoria = json.loads(c.memoria_calculo_json)
            for passo in memoria.get("passos", []):
                cel = ws.cell(row=row, column=1, value=passo.get("titulo"))
                cel.font = HEADER_FONT
                cel.fill = HEADER_FILL
                if passo.get("resultado"):
                    ws.cell(row=row, column=2, value=f"Resultado: {passo['resultado']}").font = BOLD
                row += 1
                for linha in passo.get("linhas", []):
                    ws.cell(row=row, column=1, value=linha).alignment = Alignment(wrap_text=True, vertical="top")
                    row += 1
                row += 1
        _autofit(ws, [120, 40])
    if not wb.sheetnames:
        wb.create_sheet("Memória de Cálculo")
    return _to_bytes(wb)
