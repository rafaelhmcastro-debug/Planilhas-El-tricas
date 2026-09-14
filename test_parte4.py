#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Teste headless da PARTE 4: status de revisão colaborativa e projeto portável
(PDF embutido em base64).

(a) Salva um projeto com o PDF embutido em base64 e reabre, confirmando que o
    PDF é recriado corretamente a partir do base64 (mesmo conteúdo).
(b) Altera o status de um item e confirma que ele aparece corretamente nas
    colunas novas do Excel (STATUS, REVISADO POR, DATA DA REVISÃO) e na aba
    "Painel de Status".

Não abre a interface gráfica de fato (root.withdraw()), só instancia
RevisorPDFApp para reaproveitar a lógica real de salvar/abrir projeto e
exportar Excel.

Rodar com:
    python test_parte4.py
"""

import os
import json
import base64
import tempfile
import tkinter as tk

import fitz
import openpyxl

import revisor_pdf as rp
from revisor_pdf import RevisorPDFApp, STATUS_REVISAO


def _montar_app(tmp):
    root = tk.Tk()
    root.withdraw()
    app = RevisorPDFApp(root)

    pdf_path = os.path.join(tmp, "original.pdf")
    doc = fitz.open()
    doc.new_page()
    doc.save(pdf_path)
    doc.close()

    app.doc = fitz.open(pdf_path)
    app.pdf_path = pdf_path
    app.dados_documento = {
        "projeto": "Projeto X", "codigo_doc": "DOC-9",
        "elaborado_por": "Ana", "verificado_por": "Bruno",
    }
    app.comentarios = [
        {
            "tipo": "comentario", "pagina": 0, "x": 5, "y": 5,
            "texto": "erro de teste", "classificacao": rp.CLASSIFICACOES[0],
            "responsavel_erro": "Carlos", "data": "01/01/2026",
            "status": "Pendente",
        },
    ]
    return root, app, pdf_path


def test_projeto_portavel_pdf_base64():
    with tempfile.TemporaryDirectory() as tmp:
        root, app, pdf_path = _montar_app(tmp)
        try:
            with open(pdf_path, "rb") as f:
                bytes_originais = f.read()

            destino_json = os.path.join(tmp, "projeto.json")
            rp.filedialog.asksaveasfilename = lambda *a, **k: destino_json
            rp.messagebox.showinfo = lambda *a, **k: None
            app.salvar_projeto()

            assert os.path.exists(destino_json)
            with open(destino_json, "r", encoding="utf-8") as f:
                dados = json.load(f)
            assert "pdf_base64" in dados and dados["pdf_base64"], "PDF não foi embutido em base64"
            assert base64.b64decode(dados["pdf_base64"]) == bytes_originais

            # simula reabrir o projeto num "outro computador": PDF original não existe mais lá
            app.doc.close()
            os.remove(pdf_path)

            rp.filedialog.askopenfilename = lambda *a, **k: destino_json
            app.doc = None
            app.comentarios = []
            app.abrir_projeto()

            assert app.doc is not None, "abrir_projeto deveria ter recriado o PDF a partir do base64"
            assert len(app.comentarios) == 1
            with open(app.pdf_path, "rb") as f:
                assert f.read() == bytes_originais, "conteúdo do PDF recriado não bate com o original"
        finally:
            if app.doc is not None:
                app.doc.close()
            root.destroy()


def test_status_no_excel():
    with tempfile.TemporaryDirectory() as tmp:
        root, app, pdf_path = _montar_app(tmp)
        try:
            app.revisor_atual = "Diana"
            app.comentarios[0]["status"] = "Corrigido"
            app.comentarios[0]["revisado_por"] = "Diana"
            app.comentarios[0]["data_revisao"] = "05/01/2026"

            destino_xlsx = os.path.join(tmp, "saida.xlsx")
            rp.filedialog.asksaveasfilename = lambda *a, **k: destino_xlsx
            rp.messagebox.showinfo = lambda *a, **k: None
            app.exportar_excel()

            wb = openpyxl.load_workbook(destino_xlsx)
            ws = wb["Comentários"]
            cabecalho = [ws.cell(row=1, column=j).value for j in range(1, 12)]
            assert cabecalho[-3:] == ["STATUS", "REVISADO POR", "DATA DA REVISÃO"], cabecalho

            linha1 = [ws.cell(row=2, column=j).value for j in range(1, 12)]
            assert linha1[-3:] == ["Corrigido", "Diana", "05/01/2026"], linha1

            ws3 = wb["Painel de Status"]
            valores = {}
            r = 2
            while ws3.cell(row=r, column=1).value not in (None, "TOTAL"):
                valores[ws3.cell(row=r, column=1).value] = ws3.cell(row=r, column=2).value
                r += 1
            assert valores["Corrigido"] == 1
            assert valores["Pendente"] == 0
            assert ws3.cell(row=r, column=2).value == 1  # TOTAL
        finally:
            if app.doc is not None:
                app.doc.close()
            root.destroy()


if __name__ == "__main__":
    test_projeto_portavel_pdf_base64()
    test_status_no_excel()
    print("OK: projeto portável (PDF em base64) e status refletido no Excel/Painel de Status.")
