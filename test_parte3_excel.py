#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Teste headless da PARTE 3: nova estrutura da planilha Excel.

Confere que a aba "Comentários" sai com exatamente as 8 colunas pedidas, na
ordem certa, uma linha por marcação (com o carimbo entrando com CODIGO ERRO /
QUANTIDADE / TIPO DE ERRO em branco), e que a aba "Resumo" agrupa por CODIGO
ERRO somando a coluna QUANTIDADE corretamente.

Não abre a interface gráfica: instancia RevisorPDFApp sobre uma tk.Tk() só
para reaproveitar a lógica de _linhas_excel()/exportar_excel() (Tk não exige
display para ser instanciado no Windows, só para ser mostrado).

Rodar com:
    python test_parte3_excel.py
"""

import os
import tempfile
import tkinter as tk

import openpyxl

from revisor_pdf import RevisorPDFApp, CODIGOS_ERRO, CLASSIFICACOES


COLUNAS_ESPERADAS = [
    "PROJETO", "DATA", "CODIGO ERRO", "QUANTIDADE", "TIPO DE ERRO",
    "RESPONSÁVEL PELO ERRO", "RESPONSÁVEL PELA VERIFICAÇÃO", "NÚMERO DO DOCUMENTO",
]


def _montar_app_de_teste():
    root = tk.Tk()
    root.withdraw()
    app = RevisorPDFApp(root)

    app.dados_documento = {
        "projeto": "Projeto Teste",
        "codigo_doc": "DOC-001",
        "elaborado_por": "Maria",
        "verificado_por": "João",
        "data_elaboracao": "01/01/2026",
        "data_verificacao": "05/01/2026",
    }

    cat1, cat2 = CLASSIFICACOES[0], CLASSIFICACOES[1]
    app.comentarios = [
        {
            "tipo": "comentario", "pagina": 0, "x": 10, "y": 10,
            "texto": "erro 1", "classificacao": cat1, "data": "10/01/2026",
        },
        {
            "tipo": "retangulo", "pagina": 0, "x0": 0, "y0": 0, "x1": 10, "y1": 10,
            "texto": "erro 2", "classificacao": cat1, "data": "11/01/2026",
        },
        {
            "tipo": "elipse", "pagina": 0, "x0": 0, "y0": 0, "x1": 10, "y1": 10,
            "texto": "erro 3", "classificacao": cat2, "data": "12/01/2026",
        },
        {
            "tipo": "carimbo", "pagina": 0, "x0": 0, "y0": 0, "x1": 10, "y1": 10,
            "texto": "Carimbo de verificação - Verificador: Carlos - Data: 13/01/2026",
            "classificacao": "", "verificador": "Carlos", "data_carimbo": "13/01/2026",
            "data": "13/01/2026",
        },
    ]
    return root, app


def test_colunas_comentarios_na_ordem_certa():
    root, app = _montar_app_de_teste()
    try:
        with tempfile.TemporaryDirectory() as tmp:
            destino = os.path.join(tmp, "saida.xlsx")

            # reproduz exportar_excel() sem passar pelo filedialog
            import revisor_pdf as rp
            wb_original_save = rp.openpyxl.Workbook.save

            caminho_capturado = {}

            def fake_asksaveasfilename(*a, **k):
                return destino

            rp.filedialog.asksaveasfilename = fake_asksaveasfilename
            rp.messagebox.showinfo = lambda *a, **k: None
            app.exportar_excel()

            assert os.path.exists(destino), "arquivo Excel não foi gerado"

            wb = openpyxl.load_workbook(destino)
            ws = wb["Comentários"]

            cabecalho = [ws.cell(row=1, column=j).value for j in range(1, 9)]
            assert cabecalho == COLUNAS_ESPERADAS, f"cabeçalho errado: {cabecalho}"

            # 4 marcações -> 4 linhas de dados (linhas 2 a 5)
            assert ws.cell(row=2, column=1).value == "Projeto Teste"
            assert ws.cell(row=2, column=3).value == CODIGOS_ERRO[CLASSIFICACOES[0]]
            assert ws.cell(row=2, column=4).value == 1
            assert ws.cell(row=2, column=5).value == CLASSIFICACOES[0]
            # RESPONSÁVEL PELO ERRO (Parte 6, item 13): sempre "Elaborado por" do
            # documento, não é mais um campo por item.
            assert ws.cell(row=2, column=6).value == "Maria"
            assert ws.cell(row=2, column=7).value == "João"
            assert ws.cell(row=2, column=8).value == "DOC-001"

            # linha do carimbo (linha 5): CODIGO ERRO / QUANTIDADE / TIPO DE ERRO em branco,
            # mas RESPONSÁVEL PELO ERRO ainda é preenchido com "Elaborado por" (regra global)
            linha_carimbo = [ws.cell(row=5, column=j).value for j in range(1, 9)]
            assert linha_carimbo[2] in (None, ""), "CODIGO ERRO do carimbo deveria estar em branco"
            assert linha_carimbo[3] in (None, ""), "QUANTIDADE do carimbo deveria estar em branco"
            assert linha_carimbo[4] in (None, ""), "TIPO DE ERRO do carimbo deveria estar em branco"
            assert linha_carimbo[5] == "Maria", "RESPONSÁVEL PELO ERRO deveria ser o Elaborado por do documento"
            assert linha_carimbo[6] == "Carlos", "RESPONSÁVEL PELA VERIFICAÇÃO do carimbo deveria ser o verificador do carimbo"

            # ---- aba Resumo: soma por código, ignorando o carimbo ---- #
            ws2 = wb["Resumo"]
            cabecalho2 = [ws2.cell(row=1, column=j).value for j in range(1, 4)]
            assert cabecalho2 == ["CODIGO ERRO", "TIPO DE ERRO", "QUANTIDADE"]

            somas = {}
            linha = 2
            while ws2.cell(row=linha, column=2).value not in (None, "TOTAL"):
                codigo = ws2.cell(row=linha, column=1).value
                qtd = ws2.cell(row=linha, column=3).value
                somas[codigo] = qtd
                linha += 1

            assert somas[CODIGOS_ERRO[CLASSIFICACOES[0]]] == 2, "cat1 deveria somar 2 ocorrências"
            assert somas[CODIGOS_ERRO[CLASSIFICACOES[1]]] == 1, "cat2 deveria somar 1 ocorrência"

            # linha TOTAL logo após a última categoria com dados relevantes
            total_valor = ws2.cell(row=linha, column=3).value
            assert total_valor == 3, f"TOTAL deveria ser 3 (excluindo o carimbo), obtido {total_valor}"
    finally:
        root.destroy()


if __name__ == "__main__":
    test_colunas_comentarios_na_ordem_certa()
    print("OK: 8 colunas na ordem certa, carimbo tratado à parte, Resumo soma QUANTIDADE por código.")
